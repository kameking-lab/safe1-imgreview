# -*- coding: utf-8 -*-
"""
build_v2_pptx.py — (株)博展向け 最終版 hakuten_jirei_v2.pptx を生成（新規・非破壊）。

RULES_V2 ②:
- 表紙 ＋ 15事例×2スライド（写真スライド＋項目スライド）。
- 写真スライド：base/OpenAI/Google の3枚を横並びで大きく配置。AIモデル名ラベルは載せない。
- 題名：各事例の創作タイトル（cases_v2/*.md）。表紙タイトルも内容に合わせ刷新。
- 項目：発生事象／原因概要（箇条書き）／対応（想定・箇条書き）／対策概要（箇条書き）。
- 監修：金田 義太（登録第4840号）を全スライドに。
- 出典は「参考資料」として控えめに（各事例スライド隅に小さく）。
- AI/創作注記は表紙に1行のみ。

本文は cases_v2/{N}.md をパースして取り込む（捏造なし）。
写真は photos_v16/{N}/{base,openai,google}.png。
テンプレ生成ロジック build_pptx.py のヘルパー・配色・フォント・フッターを再利用（読込のみ・非破壊）。
出力は新規ファイル hakuten_jirei_v2.pptx（既存テンプレ hakuten_jirei_cases.pptx は上書きしない）。
"""
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

from build_pptx import (
    set_font, add_text, add_footer, blank_slide, fit_cover,
    JP, BLACK, RED, GRAYL, FOOT, WHITE, DARK, REDDK,
    SW, SH,
)

BASE = r"C:\Users\kanet\20260522\safe1"
CASES_DIR = os.path.join(BASE, "cases_v2")
PHOTOS = os.path.join(BASE, "photos_v16")
OUT = os.path.join(BASE, "hakuten_jirei_v2.pptx")

TODAY = "2026年6月15日"
SUPERVISOR = "監修：金田 義太（労働安全コンサルタント 登録第4840号）"
RGB_999 = RGBColor(0x99, 0x99, 0x99)

NS = ["N%02d" % i for i in range(1, 16)]
PHOTO_FILES = ["base.png", "openai.png", "google.png"]


def parse_case(n):
    """cases_v2/{n}.md をパースして dict を返す。"""
    path = os.path.join(CASES_DIR, n + ".md")
    with open(path, encoding="utf-8") as f:
        txt = f.read()
    # ヘッダ行から カテゴリ / 事故の型
    cat, kind = "", ""
    m = re.search(r"カテゴリ：([^／/]+)[／/].*?事故の型：([^\n／/]+)", txt)
    if m:
        cat = m.group(1).strip()
        kind = m.group(2).strip()

    # セクション分割（## 見出し ごと）
    sections = {}
    cur = None
    buf = []
    for line in txt.splitlines():
        hm = re.match(r"^##\s+(.*)$", line)
        if hm:
            if cur is not None:
                sections[cur] = buf
            cur = hm.group(1).strip()
            buf = []
        elif cur is not None:
            buf.append(line)
    if cur is not None:
        sections[cur] = buf

    def sec(name_prefix):
        for k, v in sections.items():
            if k.startswith(name_prefix):
                return v
        return []

    def bullets(name_prefix):
        out = []
        for line in sec(name_prefix):
            s = line.strip()
            if s.startswith("- "):
                out.append(s[2:].strip())
        return out

    def paragraph(name_prefix):
        out = []
        for line in sec(name_prefix):
            s = line.strip()
            if s and not s.startswith("-"):
                out.append(s)
        return " ".join(out)

    title = paragraph("創作タイトル").strip()

    # 参考資料：ラベル行（- ...）＋ URL 行
    ref_label, ref_url = "", ""
    for line in sec("参考資料"):
        s = line.strip()
        if s.startswith("- ") and not ref_label:
            ref_label = s[2:].strip()
        um = re.search(r"https?://\S+", s)
        if um and not ref_url:
            ref_url = um.group(0).strip()

    return dict(
        n=n, cat=cat, kind=kind, title=title,
        event=paragraph("発生事象"),
        cause=bullets("原因概要"),
        resp=bullets("対応"),
        meas=bullets("対策概要"),
        ref_label=ref_label, ref_url=ref_url,
    )


def supervisor_line(slide):
    add_text(slide, Inches(0.5), Inches(1.16), Inches(12.3), Inches(0.28),
             [{"runs": [(SUPERVISOR, dict(name=JP, size=11, color=FOOT))]}],
             anchor=MSO_ANCHOR.MIDDLE)


def case_title(slide, c):
    txt = "事故事例 %s　%s" % (c["n"], c["title"])
    add_text(slide, Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.92),
             [{"runs": [(txt, dict(name=JP, size=22, bold=True, color=BLACK))],
               "line_spacing": 1.0}])


# ---- 写真スライド レイアウト（3枚横並び） ----
IMG_RATIO = 1.40          # W/H（歪み無し・中央クロップで統一）
IMG_W = 3.85
IMG_H = IMG_W / IMG_RATIO  # 2.75
GAP = 0.39
ROW_TOP = 2.85
COLS = [0.50, 0.50 + IMG_W + GAP, 0.50 + 2 * (IMG_W + GAP)]  # 0.50 / 4.74 / 8.98


def photo_slide(prs, c, page_no):
    s = blank_slide(prs)
    case_title(s, c)
    supervisor_line(s)
    # サブ見出し（カテゴリ・事故の型／控えめ・モデル名なし）
    sub = "%s／%s" % (c["cat"], c["kind"])
    add_text(s, Inches(0.5), Inches(1.46), Inches(12.3), Inches(0.30),
             [{"runs": [(sub, dict(name=JP, size=12, color=DARK))]}])
    pdir = os.path.join(PHOTOS, c["n"])
    for i, fn in enumerate(PHOTO_FILES):
        fp = os.path.join(pdir, fn)
        fit = fit_cover(fp, IMG_RATIO)
        pic = s.shapes.add_picture(fit, Inches(COLS[i]), Inches(ROW_TOP),
                                   Inches(IMG_W), Inches(IMG_H))
        pic.line.color.rgb = GRAYL
        pic.line.width = Pt(0.75)
    # 控えめキャプション（モデル名・AI注記なし）
    add_text(s, Inches(0.5), Inches(ROW_TOP + IMG_H + 0.12), Inches(12.3), Inches(0.30),
             [{"runs": [("※画像は選抜用の3案を併載（同一事故の再現イメージ）",
                         dict(name=JP, size=10, color=FOOT))]}],
             align=PP_ALIGN.CENTER)
    add_footer(s, page_no)


def item_slide(prs, c, page_no):
    s = blank_slide(prs)
    case_title(s, c)
    supervisor_line(s)
    rows = [
        ("発生事象", c["event"], False),
        ("原因概要", c["cause"], True),
        ("対応（想定）", c["resp"], True),
        ("対策概要", c["meas"], True),
    ]
    L = Inches(0.6)
    T = Inches(1.58)
    W = Inches(12.13)
    label_w = Inches(2.25)
    heights = [1.18, 1.10, 0.98, 1.46]
    tbl = s.shapes.add_table(len(rows), 2, L, T, W, Inches(sum(heights))).table
    tbl.first_row = False
    tbl.horz_banding = False
    tbl.columns[0].width = label_w
    tbl.columns[1].width = Emu(int(W) - int(label_w))
    for ri, (lab, val, is_list) in enumerate(rows):
        tbl.rows[ri].height = Inches(heights[ri])
        lc = tbl.cell(ri, 0)
        lc.fill.solid(); lc.fill.fore_color.rgb = GRAYL
        lc.vertical_anchor = MSO_ANCHOR.MIDDLE
        lc.margin_left = Pt(8); lc.margin_top = Pt(3); lc.margin_bottom = Pt(3)
        p = lc.text_frame.paragraphs[0]; r = p.add_run(); r.text = lab
        set_font(r, name=JP, size=12.5, bold=True, color=BLACK)
        vc = tbl.cell(ri, 1)
        vc.fill.solid(); vc.fill.fore_color.rgb = WHITE
        vc.vertical_anchor = MSO_ANCHOR.MIDDLE
        vc.margin_left = Pt(10); vc.margin_top = Pt(3); vc.margin_bottom = Pt(3)
        tf = vc.text_frame; tf.word_wrap = True
        if is_list:
            items = val if val else ["—"]
            for i, item in enumerate(items):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(2)
                r = p.add_run(); r.text = "・" + item
                set_font(r, name=JP, size=11, color=BLACK)
        else:
            p = tf.paragraphs[0]; r = p.add_run(); r.text = val
            set_font(r, name=JP, size=11, color=BLACK)
    # 参考資料（控えめ・小さく／スライド下部）
    ref = "参考資料：%s　%s" % (c["ref_label"], c["ref_url"])
    add_text(s, Inches(0.6), Inches(6.42), Inches(12.13), Inches(0.50),
             [{"runs": [(ref, dict(name=JP, size=9, color=FOOT))]}])
    add_footer(s, page_no)


def cover(prs, page_no):
    s = blank_slide(prs)
    add_text(s, Inches(0.9), Inches(1.95), Inches(11.6), Inches(2.0),
             [{"runs": [("イベント設営現場 想定事故事例集", dict(name=JP, size=40, bold=True, color=BLACK))],
               "space_after": 6},
              {"runs": [("テールゲートリフター・高所作業車（全15事例）", dict(name=JP, size=26, bold=True, color=BLACK))]}],
             align=PP_ALIGN.LEFT)
    add_text(s, Inches(0.95), Inches(4.30), Inches(8.0), Inches(0.5),
             [{"runs": [("株式会社 博展　御中", dict(name=JP, size=17, bold=True, color=BLACK))]}])
    add_text(s, Inches(0.95), Inches(4.92), Inches(8.0), Inches(0.4),
             [{"runs": [("作成日：" + TODAY, dict(name=JP, size=12, color=FOOT))]}])
    add_text(s, Inches(0.95), Inches(5.35), Inches(10.5), Inches(0.4),
             [{"runs": [(SUPERVISOR, dict(name=JP, size=11, color=FOOT))]}])
    # AI/創作注記（表紙に1行のみ）
    add_text(s, Inches(0.95), Inches(5.92), Inches(11.6), Inches(0.6),
             [{"runs": [("※本資料は、公的災害事例の機序を参考に、イベント設営現場で起こり得る事故を生成AI画像で再現した創作（安全教育用）事例集です。",
                         dict(name=JP, size=10.5, color=REDDK))]}])
    # 参考資料の出所（控えめ）
    add_text(s, Inches(0.95), Inches(6.45), Inches(11.6), Inches(0.4),
             [{"runs": [("参考資料：厚生労働省 職場のあんぜんサイト／建設荷役車両安全技術協会（各事例に参考URLを記載）",
                         dict(name=JP, size=9.5, color=FOOT))]}])
    add_footer(s, page_no)


def main():
    cases = [parse_case(n) for n in NS]
    # 簡易検証（空欄が無いか）
    for c in cases:
        assert c["title"], "no title " + c["n"]
        assert c["event"], "no event " + c["n"]
        assert c["cause"] and c["resp"] and c["meas"], "empty bullets " + c["n"]
        assert c["ref_url"].startswith("http"), "no url " + c["n"]
        for fn in PHOTO_FILES:
            assert os.path.exists(os.path.join(PHOTOS, c["n"], fn)), "missing photo " + c["n"] + " " + fn

    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    page = 0

    def npg():
        nonlocal page
        page += 1
        return page

    cover(prs, npg())
    for c in cases:
        photo_slide(prs, c, npg())
        item_slide(prs, c, npg())

    prs.save(OUT)
    print("SAVED", os.path.basename(OUT), "slides=", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
