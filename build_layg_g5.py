# -*- coding: utf-8 -*-
"""
build_layg_g5.py — G5: 合成見本スライドを既存 sample_layers_v18.pptx に追加（非破壊で開いて末尾に2枚追記）。
2シチュ：(1) TGL墜落  (2) TGL荷崩れ・下敷き。
背景を全面に敷き、その上に荷物/作業員パーツを「事故に見える位置・サイズ・回転角」で
それぞれ個別の画像オブジェクトとして配置（1枚に焼き込まない＝移動・回転・拡縮可能）。
パーツは透明余白を持つため、各パーツのアルファ境界で切り出した新ファイル(*_c.png)を作って配置。
既存ファイル(layers_v18/・G4スライド)は破壊しない。Googleのみ生成素材・OpenAI不使用。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

BASE = r"C:\Users\kanet\20260522\safe1"
LV = os.path.join(BASE, "layers_v18")
OUT = os.path.join(BASE, "sample_layers_v18.pptx")

JP = "游ゴシック"
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x00, 0x00)
GRAY = RGBColor(0x44, 0x44, 0x44)

SW_IN, SH_IN = 13.333, 7.5


def set_font(run, name=JP, size=14, bold=False, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "ja-JP")
    rPr.set("altLang", "en-US")
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", name)


def add_text(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, wrap=True):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]
    p.alignment = align
    for txt, opt in runs:
        r = p.add_run(); r.text = txt
        set_font(r, **opt)
    return tb


def banner(slide, l, t, w, h, fill=WHITE, alpha=22):
    """半透明の白帯（テキスト可読性確保）。alpha=透過率%。"""
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.line.fill.background()
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    # 透明度を設定
    srgb = sp.fill.fore_color._xFill.find(qn("a:srgbClr"))
    a = srgb.makeelement(qn("a:alpha"), {"val": str(int((100 - alpha) * 1000))})
    srgb.append(a)
    sp.shadow.inherit = False
    return sp


def crop_to_alpha(src, dst, pad=6):
    """透明余白を切り落とした新PNGを作成（非破壊・新ファイル名）。アスペクト比を返す。"""
    with Image.open(src) as im:
        im = im.convert("RGBA")
        w, h = im.size
        bb = im.split()[3].getbbox() or (0, 0, w, h)
        l = max(0, bb[0] - pad); t = max(0, bb[1] - pad)
        r = min(w, bb[2] + pad); b = min(h, bb[3] + pad)
        cr = im.crop((l, t, r, b))
        if not os.path.exists(dst):
            cr.save(dst)
        return cr.width / cr.height


def place(slide, path, aspect, cx, cy, h_in, rot=0):
    """中心(cx,cy)・高さh_in・回転rot度 で個別画像オブジェクトとして配置。"""
    bh = h_in
    bw = bh * aspect
    pic = slide.shapes.add_picture(path, Inches(cx - bw / 2), Inches(cy - bh / 2),
                                   Inches(bw), Inches(bh))
    if rot:
        pic.rotation = rot
    return pic


def comp_slide(prs, bg_path, title, subtitle, parts, note):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # 背景を全面に
    slide.shapes.add_picture(bg_path, 0, 0, Inches(SW_IN), Inches(SH_IN))
    # パーツを個別オブジェクトで配置
    for p in parts:
        place(slide, p["path"], p["aspect"], p["cx"], p["cy"], p["h"], p.get("rot", 0))
    # タイトル帯（上）
    banner(slide, 0, 0, SW_IN, 0.78, fill=WHITE, alpha=20)
    add_text(slide, 0.4, 0.06, SW_IN - 0.8, 0.42,
             [(title, dict(size=20, bold=True, color=BLACK))],
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.4, 0.46, SW_IN - 0.8, 0.3,
             [(subtitle, dict(size=11, color=GRAY))],
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    # 注記帯（下）
    banner(slide, 0, SH_IN - 0.5, SW_IN, 0.5, fill=WHITE, alpha=20)
    add_text(slide, 0.4, SH_IN - 0.48, SW_IN - 0.8, 0.44,
             [("※ これはレイヤー合成のたたき台。各部品は移動・回転・拡縮可能な個別画像オブジェクト。最終調整はPowerPoint上で。",
               dict(size=10, color=GRAY))],
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    slide.notes_slide.notes_text_frame.text = note
    return slide


def main():
    obj = os.path.join(LV, "obj")
    wk = os.path.join(LV, "worker")
    bg = os.path.join(LV, "bg")

    # 使用パーツをアルファ境界で切り出し（新ファイル *_c.png）
    a_tip = crop_to_alpha(os.path.join(obj, "obj_rollcage_tipping.png"),
                          os.path.join(obj, "obj_rollcage_tipping_c.png"))
    a_fall = crop_to_alpha(os.path.join(wk, "worker_fallen_back.png"),
                           os.path.join(wk, "worker_fallen_back_c.png"))
    a_startled = crop_to_alpha(os.path.join(wk, "worker_startled.png"),
                               os.path.join(wk, "worker_startled_c.png"))
    a_coll = crop_to_alpha(os.path.join(obj, "obj_rollcage_collapsing.png"),
                           os.path.join(obj, "obj_rollcage_collapsing_c.png"))
    a_panel = crop_to_alpha(os.path.join(obj, "obj_panelcart_sliding.png"),
                            os.path.join(obj, "obj_panelcart_sliding_c.png"))
    a_foot = crop_to_alpha(os.path.join(wk, "worker_foot_crouch.png"),
                           os.path.join(wk, "worker_foot_crouch_c.png"))
    a_hand = crop_to_alpha(os.path.join(wk, "worker_hand_pain.png"),
                           os.path.join(wk, "worker_hand_pain_c.png"))

    prs = Presentation(OUT)  # 既存（G4スライド入り）を開いて末尾に追記＝非破壊
    n0 = len(prs.slides._sldIdLst)

    # シチュ1：TGL墜落（昇降板付近でカゴ車が転倒、作業員が後方へ墜落）
    comp_slide(
        prs, os.path.join(bg, "bg_warehouse1.png"),
        "合成見本① TGL墜落（テールゲートリフターからの転落）",
        "背景=bg_warehouse1／カゴ車転倒(obj_rollcage_tipping)＋驚き後傾(worker_startled)＋転落(worker_fallen_back)。各パーツは個別配置・回転済み。",
        [
            {"path": os.path.join(obj, "obj_rollcage_tipping_c.png"), "aspect": a_tip,
             "cx": 8.7, "cy": 4.35, "h": 3.0, "rot": 22},
            {"path": os.path.join(wk, "worker_startled_c.png"), "aspect": a_startled,
             "cx": 6.0, "cy": 4.15, "h": 3.0, "rot": -10},
            {"path": os.path.join(wk, "worker_fallen_back_c.png"), "aspect": a_fall,
             "cx": 3.1, "cy": 5.7, "h": 2.4, "rot": 6},
        ],
        ("合成見本①（G5）TGL墜落。Googleのみ生成素材（gemini-3-pro-image-preview）・OpenAI不使用。"
         "背景の上にカゴ車転倒・作業員(驚き/転落)を個別の移動・回転・拡縮可能な画像オブジェクトとして配置。"
         "位置/サイズ/回転はパワポ上で微調整する前提のたたき台。"),
    )

    # シチュ2：TGL荷崩れ・下敷き（積荷が崩れ、作業員が足を負傷してしゃがむ／パネル台車も滑落）
    comp_slide(
        prs, os.path.join(bg, "bg_warehouse2.png"),
        "合成見本② TGL荷崩れ・下敷き（積荷の崩落）",
        "背景=bg_warehouse2／カゴ車崩れ(obj_rollcage_collapsing)＋パネル台車滑落(obj_panelcart_sliding)＋足負傷(worker_foot_crouch)＋手負傷(worker_hand_pain)。",
        [
            {"path": os.path.join(obj, "obj_rollcage_collapsing_c.png"), "aspect": a_coll,
             "cx": 7.4, "cy": 4.2, "h": 3.6, "rot": 10},
            {"path": os.path.join(obj, "obj_panelcart_sliding_c.png"), "aspect": a_panel,
             "cx": 10.6, "cy": 5.3, "h": 1.7, "rot": -14},
            {"path": os.path.join(wk, "worker_foot_crouch_c.png"), "aspect": a_foot,
             "cx": 3.5, "cy": 4.9, "h": 3.2, "rot": 0},
            {"path": os.path.join(wk, "worker_hand_pain_c.png"), "aspect": a_hand,
             "cx": 5.5, "cy": 4.7, "h": 3.3, "rot": -6},
        ],
        ("合成見本②（G5）TGL荷崩れ・下敷き。Googleのみ生成素材（gemini-3-pro-image-preview）・OpenAI不使用。"
         "背景の上にカゴ車崩れ・パネル台車滑落・作業員(足/手の負傷)を個別の移動・回転・拡縮可能な画像オブジェクトとして配置。"
         "位置/サイズ/回転はパワポ上で微調整する前提のたたき台。"),
    )

    prs.save(OUT)
    print("SAVED", OUT, "slides:", n0, "->", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
