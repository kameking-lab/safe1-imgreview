# -*- coding: utf-8 -*-
"""B4: あんぜんサイトDBで(1)出典URL補完 と (2)件数積み増し。

目的:
  (1) 出典URL補完 — 既存の JNIOSH 抽出全件(TGL/AERIAL)に source_url を付与し、
      B5 の「出典URL全件必須」を満たす。JNIOSH 由来 = 公表報告書ページURL。
  (2) 件数積み増し — JNIOSH整形CSVは死亡=〜2018/死傷=〜2017で打ち切り。
      あんぜんサイト公開DB(Excel)は死亡1991-2023・死傷〜2021を収録するため、
      JNIOSH に無い新しい年(死亡 R01-R05=2019-2023 / 死傷 H30-R03=2018-2021)を
      ダウンロード済みファイルから取り込み、同一の TGL/AERIAL 判定・重複排除を適用して
      既存抽出に対し新規分のみ追加する。各件に個別出典URL(xlsxファイルURL+ID)を付与。

入力(DL済・data/anzen/ 配下、追加取得しない):
  data/anzen/sib_xls/sibou_db_r0{1..5}.xlsx      死亡災害DB 2019-2023
  data/anzen/shisyo_xls/sisyou_db_{h30,h31,r01,r02,r03}_*.xlsx  死傷DB 2018-2021
既存(壊さない・読み取りのみ):
  data/jniosh/parsed/tgl_extracted.jsonl / aerial_extracted.jsonl
出力(新ファイル名・append only / 既存JNIOSH抽出は不変):
  data/jniosh/parsed/tgl_anzen_add.jsonl        あんぜんサイト由来の新規TGL
  data/jniosh/parsed/aerial_anzen_add.jsonl     あんぜんサイト由来の新規AERIAL
  data/jniosh/parsed/tgl_enriched.jsonl         全TGL(既存+新規)・全件source_url付
  data/jniosh/parsed/aerial_enriched.jsonl      全AERIAL(既存+新規)・全件source_url付
  data/jniosh/parsed/b4_summary.json            件数・内訳
再実行安全: 決定的に再生成(冪等)。重複排除は (発生年月+業種+型+状況) のハッシュ。
"""
import glob
import hashlib
import json
import os

import openpyxl

ANZEN_SIB_DIR = "data/anzen/sib_xls"
ANZEN_SHISYO_DIR = "data/anzen/shisyo_xls"
PARSED = "data/jniosh/parsed"

TGL_EXTRACTED = os.path.join(PARSED, "tgl_extracted.jsonl")
AERIAL_EXTRACTED = os.path.join(PARSED, "aerial_extracted.jsonl")

# 出典URL（公開・検証可能な一次ソース）
URL_JNIOSH = "https://www.jniosh.johas.go.jp/publication/houkoku/houkoku_2022_01.html"
URL_SIB_FND = "https://anzeninfo.mhlw.go.jp/anzen_pg/SIB_FND.html"
URL_SHISYO_FND = "https://anzeninfo.mhlw.go.jp/anzen_pgm/SHISYO_FND.html"
ANZEN_SIB_BASE = "https://anzeninfo.mhlw.go.jp/anzen/sib_xls/"
ANZEN_SHISYO_BASE = "https://anzeninfo.mhlw.go.jp/anzen/shisyo_xls/"

# ---- 抽出キーワード（B2/B3と完全一致）----
TGL_KEYWORDS = [
    "テールゲート", "パワーゲート", "パワーリフト", "リフトゲート", "テールリフト",
    "荷台昇降", "床面昇降", "後部昇降", "昇降テール", "昇降ゲート",
    "昇降装置付", "昇降板付", "格納式パワーゲート", "格納式テールゲート",
]
AERIAL_PRIMARY = ["高所作業車", "高所作業台"]
AERIAL_KIIN_CODE = "146"
AERIAL_CONTEXT = [
    "バスケット", "作業床", "ゴンドラ", "リフト車",
    "垂直昇降", "伸縮ブーム", "屈折ブーム", "ブーム",
]


def category_of(rec):
    t = (rec.get("type") or "").strip()
    return t if t else "その他・不明"


def dedup_key(rec):
    parts = [
        rec.get("seireki", ""), rec.get("tsuki", ""),
        rec.get("gyosyu_L", ""), rec.get("gyosyu_M", ""),
        rec.get("type", ""), rec.get("jokyo", ""),
    ]
    return hashlib.md5("␟".join(parts).encode("utf-8")).hexdigest()


def match_tgl(r):
    blob = "".join([r.get("jokyo", ""), r.get("kiin_S", ""),
                    r.get("kiin_M", ""), r.get("kiin_L", "")])
    keys = [k for k in TGL_KEYWORDS if k in blob]
    return (len(keys) > 0), keys


def match_aerial(r):
    blob = "".join([r.get("jokyo", ""), r.get("kiin_S", ""),
                    r.get("kiin_M", ""), r.get("kiin_L", "")])
    keys = []
    for p in AERIAL_PRIMARY:
        if p in blob:
            keys.append(p)
    if str(r.get("kiin_S_code", "")) == AERIAL_KIIN_CODE:
        keys.append("起因物コード146")
    if "高所" in blob:
        for c in AERIAL_CONTEXT:
            if c in blob:
                keys.append(c)
    return (len(keys) > 0), keys


def s(v):
    return "" if v is None else str(v).strip()


def seireki_from_sib(fname):
    # sibou_db_rNN.xlsx -> 2018+NN (令和NN)
    base = os.path.basename(fname)
    tag = base.replace("sibou_db_", "").split(".")[0]  # rNN
    num = int(tag[1:])
    return str(2018 + num) if tag[0] == "r" else ""


def seireki_from_shisyo(fname):
    # sisyou_db_{h30|h31|r01..}_MM.xlsx
    base = os.path.basename(fname)
    tag = base.replace("sisyou_db_", "").split("_")[0]  # h30/h31/rNN
    num = int(tag[1:])
    if tag[0] == "h":
        return str(1988 + num)      # h30=2018, h31=2019
    return str(2018 + num)          # r01=2019, r02=2020, r03=2021


def parse_sib(path):
    """死亡災害DB xlsx -> 正規化レコード列。年は西暦をファイル名から付与。"""
    seireki = seireki_from_sib(path)
    base = os.path.basename(path)
    src_url = ANZEN_SIB_BASE + base
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 19:
            continue
        c0 = s(row[0])
        if c0 == "" or c0 == "ID" or not c0[0].isdigit():
            continue
        jokyo = s(row[3])
        if not jokyo:
            continue
        rec = {
            "db": "SHIBO", "src_file": base, "src_row": c0,
            "data_source": "ANZEN", "anzen_id": c0,
            "source_url": src_url,
            "seireki": seireki, "nen": "", "tsuki": s(row[1]),
            "gyosyu_L": s(row[5]), "gyosyu_M": s(row[7]),
            "kiin_L_code": s(row[11]), "kiin_L": s(row[12]),
            "kiin_M_code": s(row[13]), "kiin_M": s(row[14]),
            "kiin_S_code": s(row[15]), "kiin_S": s(row[16]),
            "type_code": s(row[17]), "type": s(row[18]),
            "jokyo": jokyo,
        }
        out.append(rec)
    wb.close()
    return out


def parse_shisyo(path):
    """死傷DB xlsx -> 正規化レコード列。年は年号/年列(無ければファイル名)。"""
    base = os.path.basename(path)
    src_url = ANZEN_SHISYO_BASE + base
    file_seireki = seireki_from_shisyo(path)
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 21:
            continue
        c0 = s(row[0])
        if c0 == "" or c0 == "ID" or not c0[0].isdigit():
            continue
        jokyo = s(row[5])
        if not jokyo:
            continue
        nengo = s(row[1])
        nen = s(row[2])
        seireki = file_seireki
        if nen.isdigit():
            if nengo == "令和":
                seireki = str(2018 + int(nen))
            elif nengo == "平成":
                seireki = str(1988 + int(nen))
        rec = {
            "db": "SHISYO", "src_file": base, "src_row": c0,
            "data_source": "ANZEN", "anzen_id": c0,
            "source_url": src_url,
            "seireki": seireki, "nen": nen, "tsuki": s(row[3]),
            "gyosyu_L": s(row[7]), "gyosyu_M": s(row[9]),
            "kiin_L_code": s(row[13]), "kiin_L": s(row[14]),
            "kiin_M_code": s(row[15]), "kiin_M": s(row[16]),
            "kiin_S_code": s(row[17]), "kiin_S": s(row[18]),
            "type_code": s(row[19]), "type": s(row[20]),
            "jokyo": jokyo,
        }
        out.append(rec)
    wb.close()
    return out


def load_existing(path):
    recs = []
    keys = set()
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            r = json.loads(line)
            recs.append(r)
            keys.add(r.get("dedup_key") or dedup_key(r))
    return recs, keys


def enrich_existing_url(r):
    """既存JNIOSH抽出に出典URLを付与(無ければ)。"""
    if not r.get("source_url"):
        r = dict(r)
        r["data_source"] = r.get("data_source", "JNIOSH")
        r["source_url"] = URL_JNIOSH
    return r


def main():
    # 既存読込
    tgl_existing, tgl_keys = load_existing(TGL_EXTRACTED)
    aerial_existing, aerial_keys = load_existing(AERIAL_EXTRACTED)
    print(f"existing: TGL={len(tgl_existing)} AERIAL={len(aerial_existing)}")

    # あんぜんサイト xlsx 全パース
    anzen = []
    sib_files = sorted(glob.glob(os.path.join(ANZEN_SIB_DIR, "*.xlsx")))
    shisyo_files = sorted(glob.glob(os.path.join(ANZEN_SHISYO_DIR, "*.xlsx")))
    for p in sib_files:
        recs = parse_sib(p)
        anzen.extend(recs)
        print(f"parsed {os.path.basename(p)}: {len(recs)}")
    for p in shisyo_files:
        recs = parse_shisyo(p)
        anzen.extend(recs)
        print(f"parsed {os.path.basename(p)}: {len(recs)}")
    print(f"anzen parsed total: {len(anzen)}")

    # 判定 + 重複排除(既存・群内)で新規のみ採用
    tgl_add, aerial_add = [], []
    seen_tgl = set(tgl_keys)
    seen_aerial = set(aerial_keys)
    for r in anzen:
        k = dedup_key(r)
        ok_t, keys_t = match_tgl(r)
        if ok_t and k not in seen_tgl:
            seen_tgl.add(k)
            o = dict(r)
            o["group"] = "TGL"
            o["category"] = category_of(r)
            o["matched_keywords"] = keys_t
            o["dedup_key"] = k
            tgl_add.append(o)
        ok_a, keys_a = match_aerial(r)
        if ok_a and k not in seen_aerial:
            seen_aerial.add(k)
            o = dict(r)
            o["group"] = "AERIAL"
            o["category"] = category_of(r)
            o["matched_keywords"] = keys_a
            o["dedup_key"] = k
            aerial_add.append(o)

    # 出力: 追加分
    with open(os.path.join(PARSED, "tgl_anzen_add.jsonl"), "w",
              encoding="utf-8", newline="\n") as o:
        for r in tgl_add:
            o.write(json.dumps(r, ensure_ascii=False) + "\n")
    with open(os.path.join(PARSED, "aerial_anzen_add.jsonl"), "w",
              encoding="utf-8", newline="\n") as o:
        for r in aerial_add:
            o.write(json.dumps(r, ensure_ascii=False) + "\n")

    # 出力: enriched(既存+追加、全件source_url付)
    tgl_all = [enrich_existing_url(r) for r in tgl_existing] + tgl_add
    aerial_all = [enrich_existing_url(r) for r in aerial_existing] + aerial_add
    with open(os.path.join(PARSED, "tgl_enriched.jsonl"), "w",
              encoding="utf-8", newline="\n") as o:
        for r in tgl_all:
            o.write(json.dumps(r, ensure_ascii=False) + "\n")
    with open(os.path.join(PARSED, "aerial_enriched.jsonl"), "w",
              encoding="utf-8", newline="\n") as o:
        for r in aerial_all:
            o.write(json.dumps(r, ensure_ascii=False) + "\n")

    def breakdown(recs):
        by_cat, by_year, by_src = {}, {}, {}
        url_ok = 0
        for r in recs:
            by_cat[r["category"]] = by_cat.get(r["category"], 0) + 1
            by_year[r["seireki"]] = by_year.get(r["seireki"], 0) + 1
            src = r.get("data_source", "JNIOSH")
            by_src[src] = by_src.get(src, 0) + 1
            if r.get("source_url"):
                url_ok += 1
        return {
            "total": len(recs),
            "with_source_url": url_ok,
            "by_data_source": by_src,
            "by_category": dict(sorted(by_cat.items(), key=lambda x: -x[1])),
            "by_year": dict(sorted(by_year.items())),
        }

    summary = {
        "anzen_files": {
            "sib": [os.path.basename(p) for p in sib_files],
            "shisyo": [os.path.basename(p) for p in shisyo_files],
        },
        "anzen_rows_parsed": len(anzen),
        "added_from_anzen": {"TGL": len(tgl_add), "AERIAL": len(aerial_add)},
        "TGL": {"existing": len(tgl_existing), "added": len(tgl_add),
                "total": len(tgl_all), **breakdown(tgl_all)},
        "AERIAL": {"existing": len(aerial_existing), "added": len(aerial_add),
                   "total": len(aerial_all), **breakdown(aerial_all)},
        "source_url_policy": {
            "JNIOSH": URL_JNIOSH,
            "ANZEN_SIB": URL_SIB_FND,
            "ANZEN_SHISYO": URL_SHISYO_FND,
        },
    }
    with open(os.path.join(PARSED, "b4_summary.json"), "w",
              encoding="utf-8", newline="\n") as sf:
        json.dump(summary, sf, ensure_ascii=False, indent=2)

    print(f"added: TGL+{len(tgl_add)} AERIAL+{len(aerial_add)}")
    print(f"TGL total={len(tgl_all)}  AERIAL total={len(aerial_all)}")


if __name__ == "__main__":
    main()
