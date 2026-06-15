# -*- coding: utf-8 -*-
"""build_sample_preview.py — 体裁確認用サンプル(未QA・抜粋3事例)。
本番 hakuten_jirei_cases.pptx・元テンプレは触らない。新規 sample_preview.pptx を出力。
template_spec.md の版面値に準拠。1事例=2スライド(上=写真2x2 / 下=項目表)。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image

BASE = r"C:\Users\kanet\20260522\safe1"
PV15 = os.path.join(BASE, "photos_v15")
TMP = os.path.join(BASE, "_sample_tmp"); os.makedirs(TMP, exist_ok=True)
OUT = os.path.join(BASE, "sample_preview.pptx")
FOOT_IMG = os.path.join(BASE, "images", "ref", "footer_strip_raw.png")

JP = "游ゴシック"
BLACK=RGBColor(0,0,0); FOOT=RGBColor(0x5E,0x5E,0x5E); GRAYL=RGBColor(0xD9,0xD9,0xD9)
WHITE=RGBColor(0xFF,0xFF,0xFF); REDDK=RGBColor(0xC0,0,0); GREY99=RGBColor(0x99,0x99,0x99)

CASES = [
 {"num":"0001","cat":"TGL","title":"#0001 TGL 昇降板からの墜落・転落",
  "proj":"博展 事故事例集／運送・荷役（テールゲートリフター）",
  "date":"記載なし（ヒヤリ・ハット事例集は発生年月日を非掲載）",
  "event":"一般貨物自動車運送業の作業者が、トラック荷台での作業を終えて降りようとテールゲートリフター（昇降機）に足を載せた際、荷台とリフトの段差に気づかずバランスを崩し、地面へ転落しそうになった（ヒヤリ・ハット）。",
  "cause":"テールゲートリフターが完全に上がり切っていない状態で乗り移ろうとし、荷台とリフトの段差を見落としたこと。昇降の最終確認不足。",
  "resp":"当該作業の一時中断、昇降動作・段差の点検、関係者への注意喚起、再発防止の周知（実際の事後対応は出典に記載なし＝想定・AI整理）。",
  "meas":["テールゲートリフターの昇降を最後まで確認し、荷台と同じ高さに上がり切ってから乗り移る。",
          "荷台から昇降する際は昇降設備を正しく使用する（段差をまたいで飛び降りない）。",
          "（補足・AI整理）昇降板上では手すり・周囲を確認し、無理な姿勢での乗り移りを避ける。"],
  "src":"出典：https://anzeninfo.mhlw.go.jp/hiyari/hiy_0448.html （厚労省 職場のあんぜんサイト ヒヤリ・ハット事例集／HTTP200・本文一致確認済）"},
 {"num":"0002","cat":"TGL","title":"#0002 TGL 昇降板と車体の間に足を挟まれ",
  "proj":"博展 事故事例集／運送・荷役（テールゲートリフター）",
  "date":"記載なし（ヒヤリ・ハット事例集は発生年月日を非掲載）",
  "event":"陸上貨物運送業の作業者が、商品の積み降ろし作業中、テールゲートリフターの昇降板に乗ったまま操作を行い、ゲート（昇降板）を上昇させていた際に、足の指先を車両とゲートの間に挟みそうになった（ヒヤリ・ハット）。",
  "cause":"作業者が昇降板（リフト）に乗った状態のままテールゲートリフターを操作していたこと。可動部（車両とゲートの間）に身体（足先）が入り込む位置で操作したこと。",
  "resp":"当該作業の一時中断、操作手順・操作位置の点検、注意喚起、再発防止の周知（実際の事後対応は出典に記載なし＝想定・AI整理）。",
  "meas":["作業者が昇降板に乗った状態で荷を昇降させない（可動部から離れた位置で操作する）。",
          "荷の昇降時には昇降板のストッパーを使用する。",
          "床面積が小さく高さのある積み荷は、ロープ・ラッシングベルト等で昇降板に固定する。"],
  "src":"出典：https://anzeninfo.mhlw.go.jp/hiyari/hiy_0373.html （厚労省 職場のあんぜんサイト ヒヤリ・ハット事例集／HTTP200・本文一致確認済）"},
 {"num":"0071","cat":"高所","title":"#0071 高所作業車 下がり壁と手すりの間に挟まれ",
  "proj":"博展 事故事例集／高所作業車",
  "date":"記載なし（あんぜんサイト No.100076 は発生年月日を非掲載）",
  "event":"6階建の建設工事現場の地下室で、コンクリート壁の仕上げ（Pコン埋め）のため、被災者が高所作業車に乗って単独で移動中、扉取付部の下がり壁と高所作業車の手すりとの隙間が狭く、その間に挟まれ、同日に死亡した（死亡者1人）。",
  "cause":"通り抜けようとした扉取付用開口部の寸法に対し高所作業車のクリアランスが無く、下がり壁と手すりの間に挟まれたこと。背景に、運転の特別教育未受講、作業計画・作業指揮者の未選任。",
  "resp":"被災者の救護・救急通報、高所作業車の運転停止と現場保全、狭隘箇所（下がり壁・開口部）と通行寸法の点検、作業計画・指揮者・特別教育体制の見直しと周知（実際の事後対応は出典に記載なし＝想定・AI整理）。",
  "meas":["高所作業車の運転業務には技能講習・特別教育修了者を就かせ、指名者以外が運転しないようキーを確実に保管する。",
          "高所作業車を用いる作業は、あらかじめ作業場所に適応する作業計画を定める。",
          "作業開始前に場所・手順・分担・安全配慮事項を説明・打合せし、指揮者を定めて作業を行う。運転者に技能向上教育を実施する。"],
  "src":"出典：https://anzeninfo.mhlw.go.jp/anzen_pg/sai_det.aspx?joho_no=100076 （厚労省 職場のあんぜんサイト 労働災害事例 No.100076／HTTP200・本文一致確認済）"},
]
LABELS = ["A：実写化（OpenAI gpt-image-2）","B：実写化（Google gemini-3-pro-image-preview）",
          "C：イベント設営（OpenAI gpt-image-2）","D：イベント設営（Google gemini-3-pro-image-preview）"]
FILES = ["A_openai.png","B_google.png","C_openai_event.png","D_google_event.png"]
SUP = "監修：金田 義太（労働安全コンサルタント 登録第4840号）"

prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK = prs.slide_layouts[6]

def set_ea(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin","a:ea","a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", JP)

def tb(slide,l,t,w,h,text,size,color=BLACK,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,wrap=True):
    box=slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=box.text_frame
    tf.word_wrap=wrap; tf.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,Pt(2))
    lines=text.split("\n")
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        r=p.add_run(); r.text=ln; f=r.font; f.size=Pt(size); f.bold=bold; f.color.rgb=color; f.name=JP; set_ea(r)
    return box

def footer(slide,pageno):
    tb(slide,0.33,7.06,5.5,0.34,"©Hakuten Corporation All Rights Reserved.",9,FOOT,anchor=MSO_ANCHOR.MIDDLE)
    if os.path.exists(FOOT_IMG):
        w=Inches(3.15); h=Inches(3.15*52/520)
        slide.shapes.add_picture(FOOT_IMG,Inches(8.95),Inches(7.04),width=w,height=h)
    tb(slide,12.55,7.02,0.55,0.34,str(pageno),10,FOOT,align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)

def title_sup(slide,title):
    tb(slide,0.5,0.25,12.3,0.80,"重大事故概要 "+title,27,BLACK,bold=True)
    tb(slide,0.5,1.07,12.3,0.30,SUP,11,FOOT)

def fit_cover(src,ratio,dst):
    im=Image.open(src).convert("RGB"); w,h=im.size; tr=ratio; cr=w/h
    if cr>tr: nw=int(h*tr); x=(w-nw)//2; im=im.crop((x,0,x+nw,h))
    else: nh=int(w/tr); y=(h-nh)//2; im=im.crop((0,y,w,y+nh))
    im.save(dst); return dst

def photo_slide(num,title,pageno):
    s=prs.slides.add_slide(BLANK); title_sup(s,title)
    R=1.55; PW=3.41; PH=2.20
    colx=[2.98,6.94]; lblT=[1.50,4.28]; imgT=[1.80,4.58]
    for i in range(4):
        col=i%2; row=i//2; x=colx[col]
        # label band
        lb=tb(s,x,lblT[row],PW,0.26,LABELS[i],10,BLACK,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        lb.fill.solid(); lb.fill.fore_color.rgb=WHITE; lb.line.color.rgb=GREY99; lb.line.width=Pt(0.75)
        src=os.path.join(PV15,num,FILES[i])
        if os.path.exists(src):
            dst=os.path.join(TMP,f"{num}_{i}.png"); fit_cover(src,R,dst)
            pic=s.shapes.add_picture(dst,Inches(x),Inches(imgT[row]),width=Inches(PW),height=Inches(PH))
            pic.line.color.rgb=GRAYL; pic.line.width=Pt(0.75)
    tb(s,0.5,6.86,12.3,0.26,"※写真は選抜前の4案（A/B/C/D）を並列提示しています。最終版で1案に選抜します。",9,FOOT)
    footer(s,pageno)

def table_slide(c,pageno):
    s=prs.slides.add_slide(BLANK); title_sup(s,c["title"])
    rows=[("発覚日時",c["date"],0.55),("プロジェクト名",c["proj"],0.55),
          ("発生事象",c["event"],1.15),("原因概要",c["cause"],1.05),
          ("対応（想定）",c["resp"],0.85),("対策概要",None,1.35)]
    L,T,W=Inches(0.6),Inches(1.45),Inches(12.13); labelw=Inches(2.25)
    tbl=s.shapes.add_table(len(rows),2,L,T,W,Inches(sum(r[2] for r in rows))).table
    tbl.columns[0].width=labelw; tbl.columns[1].width=W-labelw
    tbl.first_row=False; tbl.horz_banding=False
    for ri,(lab,val,hh) in enumerate(rows):
        tbl.rows[ri].height=Inches(hh)
        lc=tbl.cell(ri,0); lc.fill.solid(); lc.fill.fore_color.rgb=GRAYL
        lc.vertical_anchor=MSO_ANCHOR.MIDDLE; lc.margin_left=Pt(8)
        lp=lc.text_frame.paragraphs[0]; lr=lp.add_run(); lr.text=lab
        lr.font.size=Pt(13); lr.font.bold=True; lr.font.name=JP; lr.font.color.rgb=BLACK; set_ea(lr)
        vc=tbl.cell(ri,1); vc.fill.solid(); vc.fill.fore_color.rgb=WHITE
        vc.vertical_anchor=MSO_ANCHOR.MIDDLE; vc.margin_left=Pt(10); vc.margin_top=Pt(4); vc.margin_bottom=Pt(4)
        vtf=vc.text_frame; vtf.word_wrap=True
        if lab=="対策概要":
            for k,m in enumerate(c["meas"]):
                p=vtf.paragraphs[0] if k==0 else vtf.add_paragraph()
                p.space_after=Pt(2); r=p.add_run(); r.text="・"+m
                r.font.size=Pt(11.5); r.font.name=JP; r.font.color.rgb=BLACK; set_ea(r)
        else:
            p=vtf.paragraphs[0]; r=p.add_run(); r.text=val
            r.font.size=Pt(12); r.font.name=JP; r.font.color.rgb=BLACK; set_ea(r)
    tb(s,0.6,T.inches+sum(r[2] for r in rows)+0.06,12.13,0.5,c["src"],9.5,FOOT)
    footer(s,pageno)

# ----- cover -----
cs=prs.slides.add_slide(BLANK)
tb(cs,0.9,1.9,11.5,1.2,"重大事故概要 事例集（体裁確認用サンプル）",30,BLACK,bold=True)
tb(cs,0.9,3.0,11.5,0.5,"株式会社 博展 御中",16,BLACK)
tb(cs,0.9,3.7,11.5,1.2,"※本資料は体裁確認用サンプルです（未QA・全17事例から抜粋3事例：#0001・#0002・#0071）。\n写真は選抜前の4案（A/B/C/D）を並列提示。発覚日時は出典に記載が無い場合「記載なし」と正直表記しています。",12.5,FOOT)
tb(cs,0.9,5.2,11.5,0.4,"作成日：2026年6月15日",12,BLACK)
tb(cs,0.9,5.7,11.5,0.4,SUP,12,FOOT)
footer(cs,1)
pg=2
for c in CASES:
    photo_slide(c["num"],c["title"],pg); pg+=1
    table_slide(c,pg); pg+=1

prs.save(OUT)
print("WROTE",OUT,"slides=",len(prs.slides.__iter__.__self__._sldIdLst))
