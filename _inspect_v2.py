# -*- coding: utf-8 -*-
from pptx import Presentation
import os
p = r"C:\Users\kanet\20260522\safe1\hakuten_jirei_v2.pptx"
prs = Presentation(p)
print("slides=", len(prs.slides._sldIdLst))
for i, s in enumerate(prs.slides):
    if i>=6: 
        print("... (more)")
        break
    texts=[]
    npic=0
    for sh in s.shapes:
        if sh.shape_type==13: npic+=1
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip().replace("\n","|")[:60])
    print(f"--- slide {i+1}: pics={npic}")
    for t in texts[:4]:
        print("   ", t)
