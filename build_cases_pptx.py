# -*- coding: utf-8 -*-
"""
build_cases_pptx.py — (株)博展向け 重大事故事例集（17事例×2スライド／写真4枚2×2＋項目表）
テンプレ(hakuten_jirei_photo_v10.pptx / build_pptx.py)の版面・配色・フッターを踏襲し、
新規ファイル hakuten_jirei_cases.pptx を生成する（テンプレ非破壊・既存pptx非上書き）。

各事例の本文は cases/{番号}.md の出典照合結果（HTTP200・本文一致確認済）に基づく要約。
写真は photos_v15/{番号}/{A_openai,B_google,C_openai_event,D_google_event}.png を A/B/C/D 2×2 配置。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# テンプレ生成ロジック(build_pptx.py)のヘルパー・配色・フォントを再利用（読込のみ・非破壊）
from build_pptx import (
    set_font, add_text, add_footer, blank_slide, fit_cover,
    JP, BLACK, RED, GRAYL, FOOT, WHITE, DARK, REDDK,
    SW, SH,
)

BASE = r"C:\Users\kanet\20260522\safe1"
PHOTOS = os.path.join(BASE, "photos_v15")
OUT = os.path.join(BASE, "hakuten_jirei_cases.pptx")

TODAY = "2026年6月15日"
SUPERVISOR = "監修：金田 義太（労働安全コンサルタント 登録第4840号）"

# 写真ラベル（2段：1段目=用途、2段目=生成モデル。gen_meta.json 由来）
PHOTO_LABELS = {
    "A_openai.png":      ("A：実写化", "OpenAI gpt-image-2"),
    "B_google.png":      ("B：実写化", "Google gemini-3-pro-image-preview"),
    "C_openai_event.png": ("C：イベント設営", "OpenAI gpt-image-2"),
    "D_google_event.png": ("D：イベント設営", "Google gemini-3-pro-image-preview"),
}
PHOTO_ORDER = ["A_openai.png", "B_google.png", "C_openai_event.png", "D_google_event.png"]

# ====== 事例データ（cases/*.md の照合結果＝出典本文の要約。捏造なし） ======
# reveal は全件「日付の記載なし」（出典に発生年月日の記載がない形式）。
CASES = [
 dict(no="0001", cat="TGL（テールゲートリフター）", kind="墜落・転落",
   short="荷台からTGL（昇降板）への乗り移り時の転落",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的事例：該当あり（MHLW）",
   proj="運送・荷役（テールゲートリフター）",
   event="一般貨物自動車運送業で、トラック荷台での作業を終えて降りる際、テールゲートリフター（昇降機）に足を載せたところ、荷台とリフトの段差でバランスを崩し転落しそうになった（ヒヤリ・ハット）。",
   cause="リフトが荷台高さまで上がり切っていない状態で乗り移ろうとし、荷台とリフトの段差を見落とした（昇降の最終確認不足）。",
   resp="（AI整理：公的出典なし）作業の一時中断、昇降動作と段差の点検、関係者への注意喚起と再発防止の周知。",
   meas=["昇降を最後まで確認し、荷台と同じ高さに上がり切ってから乗り移る",
         "荷台から昇降する際は昇降設備を正しく使用し、段差を飛び降りない"],
   src="出典：厚生労働省 職場のあんぜんサイト ヒヤリ・ハット事例 hiy_0448（https://anzeninfo.mhlw.go.jp/hiyari/hiy_0448.html）",
   site="設営搬入で資材を降ろす際、昇降板が上がり切ってから乗り移り、段差を飛び降りない。",
   note=None),
 dict(no="0002", cat="TGL（テールゲートリフター）", kind="はさまれ・巻き込まれ",
   short="昇降板に乗ったままの操作で足指をはさまれ",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的事例：該当あり（MHLW）",
   proj="運送・荷役（テールゲートリフター）",
   event="陸上貨物運送業で、商品の積み降ろし中、昇降板に乗ったままテールゲートリフターを操作してゲートを上昇させた際、足の指先を車両とゲートの間に挟みそうになった（ヒヤリ・ハット）。",
   cause="昇降板（リフト）に乗った状態のまま操作し、可動部（車両とゲートの間）に足先が入る位置で操作した。",
   resp="（AI整理：公的出典なし）作業の一時中断、操作手順・操作位置の点検、関係者への注意喚起と再発防止の周知。",
   meas=["昇降板に乗った状態で荷を昇降させない（可動部から離れた位置で操作）",
         "荷の昇降時は昇降板のストッパーを使用する",
         "背の高い積荷はロープ・ラッシングベルト等で昇降板に固定する"],
   src="出典：厚生労働省 職場のあんぜんサイト ヒヤリ・ハット事例 hiy_0373（https://anzeninfo.mhlw.go.jp/hiyari/hiy_0373.html）",
   site="昇降板の操作は乗らずに地上側から行い、可動部に手足を入れない。",
   note=None),
 dict(no="0003", cat="TGL（テールゲートリフター）", kind="墜落・転落",
   short="TGL上で台車を後ろ向きに引き、ストッパーにつまずき転落",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的事例：該当あり（MHLW）",
   proj="運送・荷役（テールゲートリフター）",
   event="道路貨物運送業で、客先での荷卸し中、荷台の台車（約200kg）を後ろ向きで引きながらテールゲートリフターに載せようとした際、リフトのストッパーにつまずきバランスを崩して落ちそうになった（ヒヤリ・ハット）。",
   cause="不安定なリフト上で重量物を後ろ向き（進行方向が見えない姿勢）で引き、足元（ストッパー）の確認が不十分だった。",
   resp="（AI整理：公的出典なし）作業の一時中断、台車移動方法と足元（ストッパー）の点検、関係者への注意喚起と再発防止の周知。",
   meas=["荷台・リフト上の台車移動は進行方向が見えるよう後方から押す（後ろ向きで引かない）",
         "荷積み・荷卸し作業の安全心得（手順）を作成し徹底する"],
   src="出典：厚生労働省 職場のあんぜんサイト ヒヤリ・ハット事例 hiy_0316（https://anzeninfo.mhlw.go.jp/hiyari/hiy_0316.html）",
   site="台車は前方が見えるよう押して移動し、リフト上の足元（段差・ストッパー）を確認する。",
   note=None),
 dict(no="0017", cat="TGL（テールゲートリフター）", kind="下敷き（飛来・落下）",
   short="昇降装置への台車移動中、ストッパー不使用で台車落下",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的事例：該当あり（MHLW）",
   proj="運送・荷役（テールゲートリフター）",
   event="貨物運送事業で、トラック荷台から雑誌を積んだ大型台車（約500kg）を昇降装置（TGL）へ単独で移動中、落下防止用ストッパーを使わなかったため台車が落ちそうになり、支えようとしたが危険を感じて避難した（下敷きのヒヤリ・ハット）。",
   cause="台車を前方に押さず進行方向に引いて操作した／落下防止用ストッパー未使用／リスクアセスメント未実施・単独作業。",
   resp="（AI整理：公的出典なし）作業の一時中断、昇降装置のストッパーと台車移動方法の点検、関係者への注意喚起と再発防止の周知。",
   meas=["台車移動は引かずに前方へ押して操作する",
         "昇降装置の落下防止用ストッパーを使用する",
         "運搬作業のリスクアセスメントを実施する",
         "単独荷卸しを禁止し複数人で行い、作業指揮者を定める"],
   src="出典：厚生労働省 職場のあんぜんサイト ヒヤリ・ハット事例 hiy_0428（https://anzeninfo.mhlw.go.jp/hiyari/hiy_0428.html）",
   site="重量台車は複数人で前方へ押し、昇降装置のストッパーを必ず使用する。",
   note=None),
 dict(no="0019", cat="TGL（※非TGL・取り違え）", kind="転倒",
   short="傾斜板からカーゴ台車が滑落・将棋倒し（※非TGL）",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的事例：該当あり（※非TGL）",
   proj="運送・荷役（荷下ろし時の台車転倒）",
   event="トラック荷台からカーゴ台車を荷下ろし中、最後尾の台車が動き出して荷台後部と地面に渡した傾斜板（ななめゲート）を滑り落ち、別の台車に衝突して2台が将棋倒しに。横で支えていた作業員はとっさに飛び降り転倒しそうになった（負傷なし）。",
   cause="動き出した最後尾のカーゴ台車にストッパーが掛けられていなかった。",
   resp="（AI整理：公的出典なし）荷下ろしの一時中断、台車ストッパーと傾斜板設置状態の点検、関係者への注意喚起と再発防止の周知。",
   meas=["荷台上の台車には必ずストッパーを掛ける",
         "積載後直ちにストッパーを掛けるよう教育する",
         "積載終了後のストッパー確認を作業指示に加える"],
   src="出典：厚生労働省 職場のあんぜんサイト ヒヤリ・ハット事例 hiy_0393（https://anzeninfo.mhlw.go.jp/hiyari/hiy_0393.html）",
   site="台車・カゴ車はキャスターロックを確実にし、傾斜板上で動き出さないようにする。",
   note="※本件は起因物＝人力運搬機（カーゴ台車）で、真正なテールゲートリフター災害ではない（img_indexのTGL分類は取り違え）。実機序は傾斜板からのカーゴ台車滑落・将棋倒し。"),
 dict(no="0040", cat="高所作業車", kind="はさまれ・巻き込まれ",
   short="バケット手摺と天井クレーンレールの間に腰部をはさまれ",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（バケットと構造物のはさまれ）",
   event="高所作業車のバケットで配線作業をしようと旋回操作したとき、据付け地盤の凹凸で車体が不安定に揺れ、バケット手摺部と頭上の天井クレーンのレール部との間に腰部をはさまれた。",
   cause="据付け地盤に凹凸があり車体が不安定だった／旋回時の上方・周囲（天井クレーンレールとの離隔）の確認不足。",
   resp="（AI整理：公的出典なし）作業の一時中断、据付け地盤と安定状態の点検、天井クレーン側との離隔確認、関係者への注意喚起と再発防止の周知。",
   meas=["据付け位置を事前に計画し、地盤の凹凸のない安定した場所で使用する",
         "作業前に車体が十分安定していること（水平・地盤）を確認してから作業する",
         "頭上構造物との離隔を確保し、旋回・上昇前に上方・周囲を確認する（AI整理の補足）"],
   src="出典：建荷協（建設荷役車両安全技術協会）労働災害事例 高０００１（https://www.sacl.or.jp/case/disaster/1135）※MHLW個別ページは該当なし（要確認）",
   site="会場の天井トラス・レール下でバケットを旋回させる前に、頭上クリアランスと据付け地盤を確認する。",
   note="※公的出典は建荷協（非MHLW）。機序一致のMHLW個別ページは該当なし（要確認）。"),
 dict(no="0041", cat="高所作業車", kind="墜落・転落",
   short="通過トラックがブームに接触、反動でバケットから転落",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（ブーム接触の反動によるバケットからの転落）",
   event="高所作業車（ブーム式）で電話線引込作業中、道路を跨ぐようブームを伸ばしていたところ、空いた車線を通過しようとしたトラックの荷台上部がブームに接触し、その反動でバケット内の作業者が転落した。",
   cause="通行車線の上方にブームを跨がせたまま作業し通過車両がブームに接触、反動が作業床に伝わった／通行路確保の計画・退避手順・墜落防止措置の不備。",
   resp="（AI整理：公的出典なし）作業の一時中断、負傷者の救護と通報、作業区域の交通規制と再確保、ブーム位置と据付け状態の点検、関係者への周知。",
   meas=["駐車車両を排除し一般車の通行路を確保して計画どおり作業する",
         "やむを得ず道路を跨ぐ場合は通過時にブームを旋回退避させる",
         "作業床上では必ず安全帯（墜落制止用器具）を使用する"],
   src="出典：建荷協 労働災害事例 高０００２（https://www.sacl.or.jp/case/disaster/1212）※MHLW個別ページは該当なし（要確認）",
   site="搬入路上空にブームを張り出す作業は通行を規制し、車両通過時はブームを退避させる。",
   note="※公的出典は建荷協（非MHLW）。機序一致のMHLW個別ページは該当なし（要確認）。"),
 dict(no="0042", cat="高所作業車", kind="はさまれ・巻き込まれ",
   short="高架橋下の剥落防止網取付け中、下面と操作盤の間にはさまれ",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的事例：該当あり（MHLW No.100614）",
   proj="高所作業車（高架橋下面と操作盤のはさまれ）",
   event="鉄道高架橋の修繕で下面にコンクリート剥落防止網を取り付ける作業中、高所作業車の作業床（作業台）が突然揺れて上方へ動き、作業者が操作盤・手摺と高架橋下面（梁）との間に挟まれ、さらに上昇して胸部圧迫により死亡した。",
   cause="高架橋下面と操作盤・手摺の離隔が狭い位置で作業床を上昇・移動／技能講習未修了者の従事、作業指揮者未選任、作業計画不備、安全教育・管理不十分。",
   resp="（AI整理：公的出典なし）作業の即時中断・緊急停止と救護・通報、作動／操作系統の点検、作業計画と指揮者体制の再構築、関係者への周知。",
   meas=["高所作業車作業に技能講習修了の有資格者を配置する",
         "作業場所・機種能力を検討し適切な作業計画を策定する",
         "作業指揮者を配置し直接指揮のもとに作業する",
         "危険予知を含む安全教育と安全管理体制を整備する"],
   src="出典：厚生労働省 職場のあんぜんサイト 労働災害事例 No.100614（https://anzeninfo.mhlw.go.jp/anzen_pg/sai_det.aspx?joho_no=100614）／補助：建荷協 高０００３",
   site="頭上に構造物がある場所では上方クリアランスを確保し、作業指揮者のもと有資格者が操作する。",
   note=None),
 dict(no="0043", cat="高所作業車", kind="墜落・転落",
   short="バケットから梁へ乗り移ろうとして足が滑り墜落",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（バケットから梁への乗り移り時の足滑りによる墜落）",
   event="高所作業車（ブーム式）のバケットが建物外周の養生ネットに引っかかり、外すため梁へ移ろうとバケット手すりに足をかけ、梁上の親綱に安全帯を掛けようとした際に足が滑り、バケットから墜落した。",
   cause="バケットから梁へ乗り移る不安定姿勢をとり、親綱への掛け替え未完了のまま足が滑り墜落を制止できなかった／高所乗降の禁止徹底・二丁掛けの不備。",
   resp="（AI整理：公的出典なし）作業の即時中断、負傷者の救護と通報、引っ掛かり解消手順の見直し、据付け状態・作業床の点検、関係者への周知。",
   meas=["高所での乗降（バケットから梁等への乗り移り）は行わない",
         "やむを得ず乗降する場合は安全な状態を確保し安全帯を使用する",
         "安全帯を二丁掛けし、掛け替え時に無胴綱状態をつくらない"],
   src="出典：建荷協 イラスト災害事例 高０００４（https://www.sacl.or.jp/case/disaster/1227）※MHLW個別ページは該当なし（要確認）",
   site="バケットから構造物へ乗り移らない。引っ掛かりはバケット操作で外し、二丁掛けを徹底する。",
   note="※公的出典は建荷協（非MHLW）。機序一致のMHLW個別ページは該当なし（要確認）。"),
 dict(no="0044", cat="高所作業車", kind="墜落・転落",
   short="作業床の手すりから上半身を出し過ぎて墜落",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（作業床の手すりから上半身を出し過ぎたことによる墜落）",
   event="高所作業車（ブーム式）の作業床上で、鉄骨間柱に胴縁材を取り付けようと一人で身を乗り出し、腕を伸ばして手すりから上半身を出し過ぎたため作業床から墜落した。",
   cause="手すりから身を乗り出し上半身を出し過ぎた無理な姿勢で体勢を崩した／作業床位置の不適切・作業床内での安全帯不使用。",
   resp="（AI整理：公的出典なし）作業の即時中断、負傷者の救護と通報、作業床の設置位置（高さ・距離）の見直し、安全帯使用の再確認、関係者への周知。",
   meas=["作業床を作業しやすい位置・高さに設置し、安全を確認してから作業する",
         "高所作業では作業床内でも安全帯（墜落制止用器具）を使用する",
         "使用時の作業手順を関係者に事前周知する"],
   src="出典：建荷協 イラスト災害事例 高０００５（https://www.sacl.or.jp/case/disaster/1231）※MHLW個別ページは該当なし（要確認）",
   site="対象物へ作業床を寄せ、手すりから身を乗り出さない。作業床内でも安全帯を使用する。",
   note="※公的出典は建荷協（非MHLW）。機序一致のMHLW個別ページは該当なし（要確認）。"),
 dict(no="0046", cat="高所作業車", kind="はさまれ・巻き込まれ",
   short="作業床上昇中、操作盤フレームと天井の間に胸部をはさまれ",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（作業床上昇中の操作盤フレームと天井のはさまれ）",
   event="競技場スタンドの工事で天井に空調ダクトを設置するため高所作業車（ブーム式）の作業床を上昇させた際、天井の段差に気付かず上昇させ続け、操作盤のフレームと天井との間に胸部を挟まれた。",
   cause="天井の段差（上方障害物）に気付かず作業床を上昇させ続け、操作盤フレームと天井の離隔がなくなった／上方確認不足・余裕のない機種選定・危険予知不足。",
   resp="（AI整理：公的出典なし）作業床の即時停止・降下と救護・通報、操作系統と手順の点検、上方障害物の事前確認体制の見直し、関係者への周知。",
   meas=["操作時は後方・上方を指差し喚呼で確認し、よそ見運転をしない",
         "操作レバーの作動パターンを熟知し事前に動作確認する",
         "作業高さに余裕のある機種を選定する",
         "危険予知訓練を行い作業指揮者を配置して作業する"],
   src="出典：建荷協 イラスト災害事例 高０００７（https://www.sacl.or.jp/case/disaster/1239）※MHLW個別ページは該当なし（要確認）",
   site="天井・梁下で作業床を上昇させる前に段差・障害物を確認し、余裕のある機種を選ぶ。",
   note="※公的出典は建荷協（非MHLW）。機序一致のMHLW個別ページは該当なし（要確認）。"),
 dict(no="0050", cat="高所作業車", kind="墜落・転落",
   short="バスケットからコンテナへ乗り移ろうとして約3m墜落",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（作業床から隣地コンテナへの乗り移り時の墜落）",
   event="樹木の枝払い後、隣地のコンテナ上の枝を片付けようと高所作業車（ブーム式）のバスケットをコンテナに近づけ、乗り移ろうとしたとき約3m下の地面に墜落した。",
   cause="作業床から他所（コンテナ）へ乗り移ろうとし、安全帯を有効に使用しないまま身体を移した／コンテナ上作業にハシゴ等を用いなかった。",
   resp="（AI整理：公的出典なし）負傷者の救護と通報、作業床からの乗り移り禁止の徹底、安全帯使用の再徹底、適切な手段の確保、手順の見直しと周知。",
   meas=["バスケット上では安全帯を着用しフックを掛け、他の場所へ乗り移らない",
         "コンテナ上の枝を除去するときはハシゴ等を使用する"],
   src="出典：建荷協 イラスト災害事例 高００１１（https://www.sacl.or.jp/case/disaster/1257）※MHLW個別ページは該当なし（要確認）",
   site="作業床から他の足場・構造物へ乗り移らない。届かない範囲は機体を寄せるか別手段を用いる。",
   note="※公的出典は建荷協（非MHLW）。機序一致のMHLW個別ページは該当なし（要確認）。"),
 dict(no="0052", cat="高所作業車", kind="転倒",
   short="傾斜地で左旋回しバランスを崩して車両が転倒",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（傾斜地での旋回時の車両転倒）",
   event="作業者2人が傾斜約10〜15度の地盤に高所作業車（ブーム式）を設置し支柱（高さ約12.5m）の塗装中、左に旋回させた際に車両がバランスを崩して転倒した。",
   cause="傾斜地に設置したまま左旋回しバランスを崩した／傾斜角の事前確認・適切な機種選定の不足、車体傾斜角警報装置を有効にしていなかった。",
   resp="（AI整理：公的出典なし）負傷者の救護と通報、傾斜地への設置可否の再確認、傾斜角測定と機種選定の徹底、安全装置の有効化、手順の見直しと周知。",
   meas=["事前に設置地盤の傾斜角を確認し適切な機種を選定する",
         "安全装置（車体傾斜角警報装置）を解除して作業しない"],
   src="出典：建荷協 イラスト災害事例 高００１３（https://www.sacl.or.jp/case/disaster/1265）※MHLW個別ページは該当なし（要確認）",
   site="屋外会場の不整地・傾斜では設置地盤を点検し、アウトリガーと水平・安全装置を確実にする。",
   note="※出典の官製「事故の型」欄は「墜落、転落」だが実機序は車両転倒（img_index分類＝転倒）。両方をそのまま記録。公的出典は建荷協（非MHLW）。"),
 dict(no="0054", cat="高所作業車", kind="墜落・転落",
   short="作業床の手すりに足をかけたダクト取付け中に約10m墜落",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（作業床の手すりに足をかけた作業中の墜落）",
   event="工場の設備工事で高所作業車に搭乗し、作業床（かご）の手すりに足をかけて天井のダクト取付け作業中、ダクト上部を固定する際にバランスを崩し約10m下のコンクリート床面に墜落した。",
   cause="作業床（かご）の手すりに足をかけた不安定姿勢で作業し固定時にバランスを崩した／安全帯を手摺に掛けず、作業床を適切な位置・高さに調整していなかった。",
   resp="（AI整理：公的出典なし）負傷者の救護と通報、墜落状況の確認、安全帯使用とフック掛けの徹底、作業床の位置・高さ調整手順の見直しと周知。",
   meas=["安全帯を着用しフックを必ず作業床（かご）の手摺に掛けてから作業する",
         "高所作業車を移動し作業床（かご）を適切な位置・高さに調整してセットする"],
   src="出典：建荷協 イラスト災害事例 高００１５（https://www.sacl.or.jp/case/disaster/1274）※MHLW個別ページは該当なし（要確認）",
   site="手すりに足をかける無理な姿勢を取らず、機体を寄せて作業床を適正な高さにセットする。",
   note="※公的出典は建荷協（非MHLW）。機序一致のMHLW個別ページは該当なし（要確認）。"),
 dict(no="0057", cat="高所作業車", kind="墜落・転落",
   short="外壁修繕中、作業床が建物に接触、確認のため身を乗り出し約20m墜落",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（外壁修繕中の作業床外への乗り出しによる墜落）",
   event="作業者とオペレーターの2名で高所作業車を使い建物外壁を修繕中、作業床の底部が建物に接触して動かなくなり、接触箇所を確認しようと作業床の外に出た作業者がバランスを崩し約20mの高さから墜落した。",
   cause="作業床が建物に接触した状態で接触箇所を確認しようと作業床の外へ身を乗り出した／フルハーネス未使用、作業床から身を乗り出さない基本動作の不徹底。",
   resp="（AI整理：公的出典なし）負傷者の救護と通報、墜落状況の確認、フルハーネス使用とフック掛けの徹底、身を乗り出さない手順の再徹底と周知。",
   meas=["高所作業ではフルハーネス型墜落制止用器具等を使用する",
         "高所作業では作業床から身を乗り出さない（接触時もいったん安全に退避させる）"],
   src="出典：建荷協 イラスト災害事例 高００１９（https://www.sacl.or.jp/case/disaster/3256）※MHLW個別ページは該当なし（要確認）",
   site="作業床が構造物に接触しても身を乗り出さず、いったん安全な位置へ退避してから確認する。",
   note="※公的出典は建荷協（非MHLW）。機序一致のMHLW個別ページは該当なし（要確認）。"),
 dict(no="0060", cat="高所作業車", kind="感電",
   short="樹木伐採後の旋回でバスケットが66kV送電線に接触、2名感電",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的：建荷協（非MHLW・MHLW個別ページ該当なし）",
   proj="高所作業車（送電線への接触による感電）",
   event="ゴルフ場改修工事で高所作業車（トラック型・伸縮ブーム）のバスケットに作業者2人が乗り樹木を伐採後、次作業へ移るため旋回したところ、バスケットが誤って66,000Vの送電線に触れ2人が感電した。",
   cause="作業車と樹木の間に66,000V送電線（推測高さ約8.4m）が通る中、送電線側へバスケットを旋回させ接触した／電力会社相談に基づく作業計画・安全離隔距離の確保不足。",
   resp="（AI整理：公的出典なし）感電者2名の救護・救急通報、電力会社への連絡と停電／防護処置の確認、近接作業の作業計画見直しと安全離隔距離の再徹底、関係者への周知。",
   meas=["事前に所管の電力会社へ相談して作業計画を立て、作業指揮者に計画に基づき指揮させる",
         "送電線の近くではクレーン等に準じて安全離隔距離を確保する（必要に応じ停電・防護・監視人配置）"],
   src="出典：建荷協 イラスト災害事例 高００２２（https://www.sacl.or.jp/case/disaster/5954）※MHLW個別ページは該当なし（要確認）",
   site="架空線・送電線の近くで作業する際は電力会社へ事前相談し、安全離隔距離と監視人を確保する。",
   note="※出典の官製「事故の型」欄は「その他」だが実機序は感電（img_index分類＝感電）。両方をそのまま記録。公的出典は建荷協（非MHLW）。"),
 dict(no="0071", cat="高所作業車", kind="はさまれ・巻き込まれ",
   short="地下室で移動中、扉の下がり壁と作業車の手すりの間にはさまれ",
   reveal="日付の記載なし（出典に発生年月日の記載なし）／公的事例：該当あり（MHLW No.100076）",
   proj="高所作業車（狭隘箇所での移動中のはさまれ）",
   event="6階建工事現場の地下室で、コンクリート壁の仕上げ（Pコン埋め）のため高所作業車に乗って単独で移動中、扉取付部の下がり壁と高所作業車の手すりとの隙間が狭く、その間に挟まれ同日に死亡した（死亡者1人）。",
   cause="通り抜けようとした扉取付用開口部の寸法に対し高所作業車の余裕（クリアランス）がなかった／特別教育未受講、作業計画・作業指揮者の未定。",
   resp="（AI整理：公的出典なし）被災者の救護・救急通報、運転停止と現場保全、狭隘箇所と通行寸法の点検、作業計画・指揮者・特別教育体制の見直しと周知。",
   meas=["運転業務は技能講習または特別教育修了者から事業者が指名した者に行わせ、キーの保管を確実にする",
         "作業場所に適応する作業計画をあらかじめ定める",
         "作業前に場所・手順・分担等を打合せし、指揮者を定めて作業する（運転者教育の実施）"],
   src="出典：厚生労働省 職場のあんぜんサイト 労働災害事例 No.100076（https://anzeninfo.mhlw.go.jp/anzen_pg/sai_det.aspx?joho_no=100076）",
   site="狭い通路・開口部を通る前に車体と壁・梁のクリアランスを確認し、指名運転者が計画に従い操作する。",
   note=None),
]

# ---- レイアウト定数（template_spec.md 3.1/3.2） ----
COL_X = {0: 2.98, 1: 6.94}        # 2列の左端(in)
IMG_W, IMG_H = 3.41, 2.20         # 画像ボックス(in) R=1.55
LBL_H = 0.40
ROW = {  # row index -> (label_top, img_top)
    0: (1.44, 1.88),
    1: (4.22, 4.64),
}

def supervisor_line(slide):
    add_text(slide, Inches(0.5), Inches(1.07), Inches(12.3), Inches(0.30),
             [{"runs": [(SUPERVISOR, dict(name=JP, size=11, color=FOOT))]}],
             anchor=MSO_ANCHOR.MIDDLE)

def case_title(slide, c):
    txt = f"重大事故概要　#{c['no']}　{c['short']}"
    add_text(slide, Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.80),
             [{"runs": [(txt, dict(name=JP, size=23, bold=True, color=BLACK))], "line_spacing": 1.0}])
    # カテゴリ・事故種類のサブ（タイトル右下の小見出し相当はフッター上に置かず監修行と分離）

def photo_cell(slide, fname, fpath, col, row):
    lx, lt = COL_X[col], ROW[row][0]
    it = ROW[row][1]
    line1, line2 = PHOTO_LABELS[fname]
    # ラベル帯（2段）
    add_text(slide, Inches(lx), Inches(lt), Inches(IMG_W), Inches(LBL_H),
             [{"runs": [(line1, dict(name=JP, size=10, bold=True, color=BLACK))],
               "align": PP_ALIGN.CENTER, "space_after": 0},
              {"runs": [(line2, dict(name=JP, size=8, color=FOOT))],
               "align": PP_ALIGN.CENTER}],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             fill=WHITE, line_color=RGB_999, line_w=0.75, wrap=True)
    # 画像（fit_cover で R=1.55 に中央クロップ＝歪みなし）
    fit = fit_cover(fpath, IMG_W / IMG_H)
    pic = slide.shapes.add_picture(fit, Inches(lx), Inches(it), Inches(IMG_W), Inches(IMG_H))
    pic.line.color.rgb = GRAYL
    pic.line.width = Pt(0.75)

from pptx.dml.color import RGBColor
RGB_999 = RGBColor(0x99, 0x99, 0x99)

def build_table(slide, c):
    rows = [
        ("発覚日時", c["reveal"], False),
        ("プロジェクト名", c["proj"], False),
        ("発生事象", c["event"], False),
        ("原因概要", c["cause"], False),
        ("対応", c["resp"], False),
        ("対策概要", c["meas"], True),
    ]
    L = Inches(0.6); T = Inches(1.4); W = Inches(12.13)
    label_w = Inches(2.25)
    heights = [0.45, 0.5, 1.15, 1.05, 0.65, 1.25]
    tbl_h = Inches(sum(heights))
    tbl = slide.shapes.add_table(len(rows), 2, L, T, W, tbl_h).table
    tbl.first_row = False; tbl.horz_banding = False
    tbl.columns[0].width = label_w
    tbl.columns[1].width = Emu(int(W) - int(label_w))
    for ri, (lab, val, is_list) in enumerate(rows):
        tbl.rows[ri].height = Inches(heights[ri])
        lc = tbl.cell(ri, 0)
        lc.fill.solid(); lc.fill.fore_color.rgb = GRAYL
        lc.vertical_anchor = MSO_ANCHOR.MIDDLE
        lc.margin_left = Pt(8); lc.margin_top = Pt(3); lc.margin_bottom = Pt(3)
        p = lc.text_frame.paragraphs[0]; r = p.add_run(); r.text = lab
        set_font(r, name=JP, size=12.5, bold=True, color=BLACK)
        vc = tbl.cell(ri, 1)
        vc.fill.solid(); vc.fill.fore_color.rgb = WHITE
        vc.vertical_anchor = MSO_ANCHOR.MIDDLE
        vc.margin_left = Pt(10); vc.margin_top = Pt(3); vc.margin_bottom = Pt(3)
        tf = vc.text_frame; tf.word_wrap = True
        if is_list:
            for i, item in enumerate(val):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(2)
                r = p.add_run(); r.text = "・" + item
                set_font(r, name=JP, size=11.5, color=BLACK)
        else:
            p = tf.paragraphs[0]; r = p.add_run(); r.text = val
            set_font(r, name=JP, size=11.5, color=BLACK)

def main():
    prs = Presentation()
    prs.slide_width = SW; prs.slide_height = SH
    page = 0
    def newpage():
        nonlocal page; page += 1; return page

    # ---- 表紙 ----
    s = blank_slide(prs)
    add_text(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.0),
      [{"runs": [("テールゲートリフター・高所作業車", dict(name=JP, size=34, bold=True, color=BLACK))], "space_after": 6},
       {"runs": [("重大事故事例集（全17事例）", dict(name=JP, size=42, bold=True, color=BLACK))]}], align=PP_ALIGN.LEFT)
    add_text(s, Inches(0.95), Inches(4.05), Inches(11.4), Inches(0.6),
      [{"runs": [("設営・搬入・撤去・会場高所作業の現場で起きる重大災害と対策／各事例に生成画像 A〜D（4案）を併載", dict(name=JP, size=14, color=DARK))]}])
    add_text(s, Inches(0.95), Inches(5.15), Inches(8.0), Inches(0.5),
      [{"runs": [("株式会社 博展　御中", dict(name=JP, size=17, bold=True, color=BLACK))]}])
    add_text(s, Inches(0.95), Inches(5.75), Inches(8.0), Inches(0.4),
      [{"runs": [("作成日：" + TODAY, dict(name=JP, size=12, color=FOOT))]}])
    add_text(s, Inches(0.95), Inches(6.18), Inches(10.5), Inches(0.4),
      [{"runs": [(SUPERVISOR, dict(name=JP, size=11, color=FOOT))]}])
    add_text(s, Inches(0.95), Inches(6.55), Inches(11.4), Inches(0.4),
      [{"runs": [("出典：厚生労働省 職場のあんぜんサイト／建設荷役車両安全技術協会（各事例ページに採用URLを明記）", dict(name=JP, size=9.5, color=FOOT))]}])
    add_footer(s, newpage())

    # ---- 各事例：写真スライド＋表スライド ----
    for c in CASES:
        pdir = os.path.join(PHOTOS, c["no"])
        # 写真スライド
        s = blank_slide(prs)
        case_title(s, c)
        supervisor_line(s)
        for idx, fname in enumerate(PHOTO_ORDER):
            col = idx % 2; row = idx // 2
            photo_cell(s, fname, os.path.join(pdir, fname), col, row)
        add_footer(s, newpage())

        # 表スライド
        s = blank_slide(prs)
        case_title(s, c)
        supervisor_line(s)
        build_table(s, c)
        # 下部注記（出典＋必要なら注記＋博展現場）
        notes = [{"runs": [(c["src"], dict(name=JP, size=9.5, color=FOOT))], "space_after": 2}]
        if c.get("note"):
            notes.append({"runs": [(c["note"], dict(name=JP, size=9, color=REDDK))], "space_after": 2})
        if c.get("site"):
            notes.append({"runs": [("博展の現場では：", dict(name=JP, size=9.5, bold=True, color=REDDK)),
                                    (c["site"], dict(name=JP, size=9.5, color=DARK))]})
        add_text(s, Inches(0.6), Inches(6.18), Inches(12.13), Inches(0.78), notes)
        add_footer(s, newpage())

    prs.save(OUT)
    print("SAVED", OUT, "slides=", len(prs.slides._sldIdLst))

if __name__ == "__main__":
    main()
