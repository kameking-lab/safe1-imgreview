# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pptx import Presentation
prs = Presentation(r"C:\Users\kanet\20260522\safe1\hakuten_jirei_v2.pptx")
print("slides=", len(prs.slides._sldIdLst))
for i, s in enumerate(prs.slides):
    npic = sum(1 for sh in s.shapes if sh.shape_type == 13)
    ntbl = sum(1 for sh in s.shapes if sh.has_table)
    title = ""
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            title = sh.text_frame.text.strip().splitlines()[0]
            break
    print(f"s{i+1:02d} pics={npic} tbl={ntbl} | {title[:50]}")
