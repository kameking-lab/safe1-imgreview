# -*- coding: utf-8 -*-
"""
build_prop.py — 商談用 提案資料 proposal_hakuten.pptx 生成（メラビアン重視・ビジュアル7割）

非破壊：既存ファイルは読むだけ。出力先は引数で指定（既定 proposal_hakuten.pptx）。
スライド構成は SLIDES に builder 関数を登録して積み上げる（P2〜P6 で順次追加）。
日本語フォントは游ゴシック（pptx 側で typeface 指定・禁則有効）。16:9。
数値・法令は RULES_PROP.md の確定素材のみ（捏造なし）。新規 AI 画像生成なし。
"""
import os, sys, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from PIL import Image

BASE = r"C:\Users\kanet\20260522\safe1"
FIGS = os.path.join(BASE, "figs")

# ---- 配色（安全色＋ブランド）------------------------------------------------
NAVY   = RGBColor(0x1F, 0x3A, 0x60)   # データ・ブランド基調
INK    = RGBColor(0x1A, 0x1A, 0x1A)   # 本文黒
GRAY   = RGBColor(0x5E, 0x5E, 0x5E)   # 補助・出典
LGRAY  = RGBColor(0x8A, 0x8A, 0x8A)
RED    = RGBColor(0xE2, 0x23, 0x1A)   # 危険
YELLOW = RGBColor(0xF2, 0xB7, 0x05)   # 注意
GREEN  = RGBColor(0x2E, 0x9E, 0x5B)   # 対策
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
PALE   = RGBColor(0xF3, 0xF5, 0xF8)   # 薄地カード
GRD    = RGBColor(0xD9, 0xD9, 0xD9)   # 表ラベル列グレー（博展テンプレ準拠）

JP  = "游ゴシック"
JPM = "游ゴシック Medium"
SW, SH = Inches(13.333), Inches(7.5)

# 確定素材（オーナー検証済・改変禁止）
SUP = "監修：金田 義太（労働安全コンサルタント 登録第4840号）"
CLIENT = "株式会社 博展　御中"
TODAY = "2026年6月21日"


# ---- 低レベル描画ヘルパ ------------------------------------------------------
def set_font(run, name=JP, size=14, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "ja-JP"); rPr.set("altLang", "en-US")
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", name)


def add_text(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             fill=None, line_color=None, line_w=None, wrap=True):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if ln.get("space_after") is not None: p.space_after = Pt(ln["space_after"])
        if ln.get("space_before") is not None: p.space_before = Pt(ln["space_before"])
        if ln.get("line_spacing"): p.line_spacing = ln["line_spacing"]
        for (txt, opts) in ln["runs"]:
            r = p.add_run(); r.text = txt
            set_font(r, **opts)
    if fill is not None:
        tb.fill.solid(); tb.fill.fore_color.rgb = fill
    else:
        tb.fill.background()
    if line_color is not None:
        tb.line.color.rgb = line_color; tb.line.width = Pt(line_w or 0.75)
    else:
        tb.line.fill.background()
    return tb


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, l, t, w, h, fill, line=None, line_w=0.0, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is not None:
        sh.line.color.rgb = line; sh.line.width = Pt(line_w or 1.0)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def circle_glyph(slide, cx, cy, d, fill, glyph, gsize, gcolor=WHITE):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx - d / 2)), Emu(int(cy - d / 2)), Emu(int(d)), Emu(int(d)))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = False
    tf.margin_top = Pt(0); tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = glyph
    set_font(r, name=JP, size=gsize, bold=True, color=gcolor)
    return sh


def add_arrow(slide, x1, y1, x2, y2, color=RED, width=3.2):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color; cn.line.width = Pt(width)
    ln = cn.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "lg", "len": "lg"}))
    return cn


def add_triangle(slide, l, t, size, mark="！"):
    sh = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, l, t, size, size)
    sh.fill.solid(); sh.fill.fore_color.rgb = YELLOW
    sh.line.color.rgb = INK; sh.line.width = Pt(1.5); sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = False
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = mark
    set_font(r, name=JP, size=14, bold=True, color=INK)
    return sh


def img_ratio(path):
    im = Image.open(path); w, h = im.size; return w / h


def place_fig(slide, path, box_l, box_t, box_w, box_h, line=None):
    """box(EMU) 内に縦横比維持で中央配置。歪みなし。枠線オプション。"""
    r = img_ratio(path)
    bw, bh = int(box_w), int(box_h)
    if bw / bh > r:
        nh = bh; nw = int(bh * r)
    else:
        nw = bw; nh = int(bw / r)
    l = int(box_l) + (bw - nw) // 2
    t = int(box_t) + (bh - nh) // 2
    pic = slide.shapes.add_picture(path, Emu(l), Emu(t), Emu(nw), Emu(nh))
    if line is not None:
        pic.line.color.rgb = line; pic.line.width = Pt(0.75)
    return pic


def tricolor_rule(slide, l, t, w, h=Inches(0.07)):
    """安全色（赤・黄・緑）の3分割細バー。"""
    seg = int(w) // 3
    for i, c in enumerate((RED, YELLOW, GREEN)):
        rect(slide, Emu(int(l) + i * seg), t, Emu(seg), h, c)


def footer(slide, page_no, note=None):
    add_text(slide, Inches(0.4), Inches(7.04), Inches(9.5), Inches(0.34),
             [{"runs": [(note or SUP, dict(name=JP, size=9, color=LGRAY))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(12.55), Inches(7.04), Inches(0.55), Inches(0.34),
             [{"runs": [(str(page_no), dict(name=JP, size=10, color=LGRAY))]}],
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def kicker(slide, text, color):
    """左上の小さな色タグ見出し。"""
    w = Inches(0.5 + 0.16 * len(text))
    add_text(slide, Inches(0.6), Inches(0.42), w, Inches(0.42),
             [{"runs": [(text, dict(name=JP, size=13, bold=True, color=WHITE))]}],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill=color)
    return Inches(0.6) + int(w)


# ============================ スライド ========================================
def slide_cover(prs, page):
    s = blank_slide(prs)
    # 左の縦帯（ネイビー）＋上部トライカラー
    rect(s, 0, 0, Inches(0.22), SH, NAVY)
    tricolor_rule(s, Inches(0.22), 0, Emu(int(SW) - int(Inches(0.22))), Inches(0.10))
    # タイトル（2行・大）
    add_text(s, Inches(1.05), Inches(1.55), Inches(11.4), Inches(1.1),
             [{"runs": [("データに基づく安全管理", dict(name=JP, size=46, bold=True, color=NAVY))]}])
    add_text(s, Inches(1.05), Inches(2.62), Inches(11.4), Inches(1.1),
             [{"runs": [("＋ ", dict(name=JP, size=46, bold=True, color=RED)),
                        ("AIによる自動化", dict(name=JP, size=46, bold=True, color=NAVY))]}])
    # サブ（価値1行）
    add_text(s, Inches(1.08), Inches(3.92), Inches(11.0), Inches(0.6),
             [{"runs": [("経験と勘に頼らない、科学的な安全教育を、速く・安く。",
                         dict(name=JP, size=18, color=GRAY))]}])
    # 区切り線
    rect(s, Inches(1.08), Inches(4.78), Inches(4.4), Inches(0.04), NAVY)
    # 宛先・監修
    add_text(s, Inches(1.05), Inches(5.30), Inches(9.0), Inches(0.6),
             [{"runs": [(CLIENT, dict(name=JP, size=22, bold=True, color=INK))]}])
    add_text(s, Inches(1.08), Inches(6.10), Inches(10.5), Inches(0.4),
             [{"runs": [(SUP, dict(name=JP, size=12, color=GRAY))]}])
    add_text(s, Inches(1.08), Inches(6.55), Inches(10.5), Inches(0.4),
             [{"runs": [("作成日：" + TODAY, dict(name=JP, size=11, color=LGRAY))]}])
    return page  # 表紙はページ番号なし


def slide_problem(prs, page):
    s = blank_slide(prs)
    kicker(s, "課題", RED)
    add_text(s, Inches(1.7), Inches(0.40), Inches(11.0), Inches(0.5),
             [{"runs": [("従来の安全教育は、ベテランの“経験頼み”",
                         dict(name=JP, size=26, bold=True, color=INK))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    # 左：大きな円（勘）
    cx, cy = int(Inches(3.2)), int(Inches(3.9))
    circle_glyph(s, cx, cy, int(Inches(2.6)), NAVY, "勘", 96)
    add_text(s, Inches(1.6), Inches(5.45), Inches(3.2), Inches(0.5),
             [{"runs": [("経験・勘に依存", dict(name=JP, size=18, bold=True, color=NAVY))]}],
             align=PP_ALIGN.CENTER)
    # 注意三角
    add_triangle(s, Inches(4.55), Inches(3.35), Inches(0.95))
    # 矢印（赤）→ 右の課題チップ3つ
    add_arrow(s, Inches(4.7), Inches(3.9), Inches(6.0), Inches(3.9), color=RED, width=4.0)
    chips = [("属人的", "教える人で内容が変わる"),
             ("再現性がない", "同じ品質を保てない"),
             ("更新が遅い", "最新の事故に追いつけない")]
    cx0, cy0, cw, ch, gap = Inches(6.25), Inches(2.05), Inches(6.5), Inches(1.05), Inches(0.30)
    for i, (head, sub) in enumerate(chips):
        ty = int(cy0) + i * (int(ch) + int(gap))
        rect(s, cx0, Emu(ty), cw, ch, PALE, line=YELLOW, line_w=2.0)
        rect(s, cx0, Emu(ty), Inches(0.14), ch, YELLOW)
        add_text(s, Emu(int(cx0) + int(Inches(0.34))), Emu(ty + int(Inches(0.12))),
                 Emu(int(cw) - int(Inches(0.5))), Inches(0.5),
                 [{"runs": [(head, dict(name=JP, size=20, bold=True, color=INK))]}])
        add_text(s, Emu(int(cx0) + int(Inches(0.34))), Emu(ty + int(Inches(0.58))),
                 Emu(int(cw) - int(Inches(0.5))), Inches(0.4),
                 [{"runs": [(sub, dict(name=JP, size=13, color=GRAY))]}])
    add_text(s, Inches(0.6), Inches(6.30), Inches(12.1), Inches(0.5),
             [{"runs": [("→ 仕組みで支える、科学的な安全教育へ。",
                         dict(name=JP, size=15, bold=True, color=NAVY))]}])
    footer(s, page)
    return page + 1


def slide_approach(prs, page):
    s = blank_slide(prs)
    kicker(s, "解決アプローチ", NAVY)
    # 矢印フロー図（P1生成・既存図を流用。図に見出しを内蔵するため大見出しは置かない）
    fig = os.path.join(FIGS, "prop_approach_flow.png")
    place_fig(s, fig, Inches(0.6), Inches(1.35), Inches(12.13), Inches(4.6))
    add_text(s, Inches(0.6), Inches(6.30), Inches(12.1), Inches(0.5),
             [{"runs": [("分析 → 危険の科学的特定 → 教材化 → AIで量産。属人化せず、速く・安く。",
                         dict(name=JP, size=14, color=GRAY))]}])
    footer(s, page)
    return page + 1


def head_bar(slide, l, t, w, text, color=NAVY):
    """チャート上の小見出しバー（色帯＋白文字）。"""
    rect(slide, l, t, w, Inches(0.46), color)
    add_text(slide, l, t, w, Inches(0.46),
             [{"runs": [(text, dict(name=JP, size=15, bold=True, color=WHITE))]}],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def stat_card(slide, l, t, w, h, accent, big, big_unit, label, sub):
    """大きな数字＋短いラベルの統計カード（安全色アクセント）。"""
    rect(slide, l, t, w, h, PALE, line=accent, line_w=2.0)
    rect(slide, l, t, Inches(0.16), h, accent)
    inx = Emu(int(l) + int(Inches(0.42)))
    inw = Emu(int(w) - int(Inches(0.62)))
    add_text(slide, inx, Emu(int(t) + int(Inches(0.10))), inw, Inches(0.86),
             [{"runs": [(big, dict(name=JP, size=46, bold=True, color=accent)),
                        (big_unit, dict(name=JP, size=18, bold=True, color=accent))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, inx, Emu(int(t) + int(Inches(0.96))), inw, Inches(0.40),
             [{"runs": [(label, dict(name=JP, size=14, bold=True, color=INK))]}])
    if sub:
        add_text(slide, inx, Emu(int(t) + int(Inches(1.33))), inw, Inches(0.32),
                 [{"runs": [(sub, dict(name=JP, size=11, color=GRAY))]}])


def slide_data1(prs, page):
    """データ分析①：規模感＋TGL/高所 型別ランキング横棒。"""
    s = blank_slide(prs)
    kx = kicker(s, "データ分析①", NAVY)
    add_text(s, Emu(int(kx) + int(Inches(0.25))), Inches(0.42), Inches(9.8), Inches(0.42),
             [{"runs": [("約", dict(name=JP, size=20, bold=True, color=INK)),
                        ("42万件", dict(name=JP, size=26, bold=True, color=RED)),
                        ("のデータから、型別の危険を特定", dict(name=JP, size=20, bold=True, color=INK))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(0.62), Inches(1.02), Inches(12.1), Inches(0.34),
             [{"runs": [("死亡DB 1991–2018 ＋ 死傷DB 2006–2017 を走査して抽出",
                         dict(name=JP, size=12, color=GRAY))]}])
    # 左右2チャート（小見出しバー付き）
    lx, rx, cw = Inches(0.62), Inches(6.92), Inches(5.78)
    head_bar(s, lx, Inches(1.46), cw, "TGL（計 1,878件）", NAVY)
    head_bar(s, rx, Inches(1.46), cw, "高所作業車（計 1,149件）", NAVY)
    place_fig(s, os.path.join(FIGS, "fig_tgl_type_rank.png"),
              lx, Inches(2.00), cw, Inches(4.10))
    place_fig(s, os.path.join(FIGS, "fig_aerial_type_rank.png"),
              rx, Inches(2.00), cw, Inches(4.10))
    # 結論バンド（安全色）
    rect(s, Inches(0.62), Inches(6.26), Inches(12.08), Inches(0.56), PALE, line=RED, line_w=1.5)
    add_text(s, Inches(0.80), Inches(6.26), Inches(11.8), Inches(0.56),
             [{"runs": [("TGL＝", dict(name=JP, size=16, bold=True, color=INK)),
                        ("はさまれ最多 30.9%", dict(name=JP, size=16, bold=True, color=RED)),
                        ("／高所＝", dict(name=JP, size=16, bold=True, color=INK)),
                        ("墜落最多 34.1%", dict(name=JP, size=16, bold=True, color=RED))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    footer(s, page)
    return page + 1


def slide_data2(prs, page):
    """データ分析②：死亡vs死傷の対比＋致死率/はさまれ過半を大きな数字で。"""
    s = blank_slide(prs)
    kx = kicker(s, "データ分析②", NAVY)
    add_text(s, Emu(int(kx) + int(Inches(0.25))), Inches(0.42), Inches(9.8), Inches(0.42),
             [{"runs": [("死亡と死傷の差 ＝ 守るべき急所", dict(name=JP, size=24, bold=True, color=INK))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    # 左：死亡vs死傷チャート
    head_bar(s, Inches(0.62), Inches(1.30), Inches(6.30), "死亡 vs 死傷（件数）", NAVY)
    place_fig(s, os.path.join(FIGS, "fig_death_vs_injury.png"),
              Inches(0.62), Inches(1.84), Inches(6.30), Inches(4.30))
    # 右：統計カード3枚
    cx, cwd = Inches(7.30), Inches(5.40)
    stat_card(s, cx, Inches(1.30), cwd, Inches(1.46), RED,
              "約31", "%", "高所作業車の致死率", "死亡358 / 全1,149件")
    stat_card(s, cx, Inches(2.92), cwd, Inches(1.46), NAVY,
              "約6", "%", "TGLの致死率", "死亡121 / 全1,878件")
    stat_card(s, cx, Inches(4.54), cwd, Inches(1.60), RED,
              "52.9", "%", "TGL死亡は「はさまれ」が過半", "死亡121件中 64件")
    # 結論バンド（安全色）
    rect(s, Inches(0.62), Inches(6.28), Inches(12.08), Inches(0.54), PALE, line=GREEN, line_w=1.5)
    add_text(s, Inches(0.80), Inches(6.28), Inches(11.8), Inches(0.54),
             [{"runs": [("高所＝", dict(name=JP, size=15, bold=True, color=INK)),
                        ("墜落で死ぬ", dict(name=JP, size=15, bold=True, color=RED)),
                        ("／TGL＝", dict(name=JP, size=15, bold=True, color=INK)),
                        ("はさまれで死ぬ", dict(name=JP, size=15, bold=True, color=RED)),
                        ("。守る急所が違う。", dict(name=JP, size=15, bold=True, color=INK))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    footer(s, page)
    return page + 1


def slide_measures(prs, page):
    """科学的対策：型別の対策をアイコン＋短句で（法令根拠は小さく・確定素材）。"""
    s = blank_slide(prs)
    kicker(s, "科学的対策", GREEN)
    add_text(s, Inches(2.05), Inches(0.42), Inches(10.6), Inches(0.42),
             [{"runs": [("急所が違えば、", dict(name=JP, size=22, bold=True, color=INK)),
                        ("打つ手", dict(name=JP, size=22, bold=True, color=GREEN)),
                        ("も変わる", dict(name=JP, size=22, bold=True, color=INK))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    # 型別の対策アイコン図（P1生成・図に見出し内蔵）
    place_fig(s, os.path.join(FIGS, "prop_measures_icons.png"),
              Inches(0.6), Inches(1.28), Inches(12.13), Inches(4.40))
    # 法令根拠（小さく・オーナー確定素材のみ）
    rect(s, Inches(0.62), Inches(5.92), Inches(12.08), Inches(0.78), PALE, line=GREEN, line_w=1.2)
    add_text(s, Inches(0.82), Inches(5.92), Inches(11.7), Inches(0.78),
             [{"runs": [("根拠（安衛則）：", dict(name=JP, size=11, bold=True, color=GREEN))],
               "space_after": 2},
              {"runs": [("フルハーネス 6.75m超 着用義務／高所作業車 作業床10m以上 技能講習・未満 特別教育／TGL特別教育 学科4h＋実技2h",
                         dict(name=JP, size=11, color=GRAY))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    footer(s, page)
    return page + 1


def slide_value(prs, page):
    """自動化の価値（売り）：教材づくりの対比＋3アイコン（属人化しない/速い/低コスト）。"""
    s = blank_slide(prs)
    kx = kicker(s, "自動化の価値", RED)
    add_text(s, Emu(int(kx) + int(Inches(0.25))), Inches(0.42), Inches(9.6), Inches(0.42),
             [{"runs": [("AIで事故事例教材を", dict(name=JP, size=22, bold=True, color=INK)),
                        ("量産", dict(name=JP, size=22, bold=True, color=RED))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    # 上：人手×日 → AI×時間 の対比図（P1生成・図に見出し内蔵）
    place_fig(s, os.path.join(FIGS, "prop_automation_compare.png"),
              Inches(0.62), Inches(1.18), Inches(12.08), Inches(2.86))
    # 下：3つの価値アイコン（P1生成）
    place_fig(s, os.path.join(FIGS, "prop_value_icons.png"),
              Inches(0.62), Inches(4.18), Inches(12.08), Inches(2.62))
    footer(s, page)
    return page + 1


def thumb_card(slide, l, t, w, h, img, caption, accent=NAVY):
    """成果物サムネイル：枠付き薄地カード＋画像（縦横比維持）＋色キャプションバー。"""
    rect(slide, l, t, w, h, PALE, line=accent, line_w=1.5)
    cap_h = Inches(0.52)
    pad = Inches(0.13)
    place_fig(slide, img,
              Emu(int(l) + int(pad)), Emu(int(t) + int(pad)),
              Emu(int(w) - 2 * int(pad)), Emu(int(h) - int(cap_h) - 2 * int(pad)),
              line=LGRAY)
    rect(slide, l, Emu(int(t) + int(h) - int(cap_h)), w, cap_h, accent)
    add_text(slide, l, Emu(int(t) + int(h) - int(cap_h)), w, cap_h,
             [{"runs": [(caption, dict(name=JP, size=13, bold=True, color=WHITE))]}],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def slide_samples(prs, page):
    """成果物サンプル：既存の学習資料/事例図を縮小し『作れます』を見せる（AI注記・小さく）。"""
    s = blank_slide(prs)
    kicker(s, "成果物サンプル", NAVY)
    add_text(s, Inches(2.55), Inches(0.42), Inches(10.1), Inches(0.42),
             [{"runs": [("こういう教材が、", dict(name=JP, size=22, bold=True, color=INK)),
                        ("すぐ作れます", dict(name=JP, size=22, bold=True, color=NAVY))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    cards = [
        ("figs/fig_qual_table.png", "資格・装備 早見表", NAVY),
        ("figs/fig_danger_points.png", "危険ポイント 図解", GREEN),
        ("photos_v16/N01/google.png", "事故事例 教材（AI再現イメージ）", RED),
    ]
    cw, ch, gap = Inches(3.95), Inches(4.50), Inches(0.28)
    x0, y0 = Inches(0.62), Inches(1.32)
    for i, (rel, cap, acc) in enumerate(cards):
        lx = Emu(int(x0) + i * (int(cw) + int(gap)))
        thumb_card(s, lx, y0, cw, ch, os.path.join(BASE, rel), cap, acc)
    # AI再現イメージ注記（小さく）
    add_text(s, Inches(0.62), Inches(6.02), Inches(12.1), Inches(0.4),
             [{"runs": [("※ 事故事例の写真はAIによる再現イメージです（実写ではありません）。",
                         dict(name=JP, size=11, color=GRAY))]}])
    footer(s, page)
    return page + 1


# ---- 指定フォーマット事故事例（博展テンプレ：上=写真／下=項目別記載）---------
def _trim(t, n):
    t = (t or "").strip()
    return t if len(t) <= n else t[:n - 1] + "…"


def parse_case(num):
    """cases_v2/<num>.md を読み、版面に流す項目を抽出（捏造なし・本文そのまま）。"""
    path = os.path.join(BASE, "cases_v2", num + ".md")
    txt = open(path, encoding="utf-8").read()
    cat = typ = ""
    m = re.search(r"カテゴリ：([^／\n]+)／事故の型：([^\n／]+)", txt)
    if m:
        cat, typ = m.group(1).strip(), m.group(2).strip()
        cat = re.sub(r"（.*?）", "", cat).strip()  # 「TGL（…）」→「TGL」
    secs, cur, buf = {}, None, []
    for line in txt.splitlines():
        if line.startswith("## "):
            if cur is not None:
                secs[cur] = buf
            cur, buf = line[3:].strip(), []
        elif cur is not None:
            buf.append(line)
    if cur is not None:
        secs[cur] = buf

    def get(key):
        for k, v in secs.items():
            if k.startswith(key):
                return v
        return []

    def joined(key):
        return " ".join(l.strip() for l in get(key) if l.strip())

    def bullets(key):
        return [l.strip()[1:].strip() for l in get(key) if l.strip().startswith("-")]

    url, src = "", ""
    for l in get("参考資料"):
        s = l.strip()
        if not src and s.startswith("-"):
            src = re.sub(r"（.*?）", "", s[1:]).strip()
        mm = re.search(r"(https?://\S+)", l)
        if mm and not url:
            url = mm.group(1)
    return dict(cat=cat, typ=typ, title=joined("創作タイトル"),
                happen=joined("発生事象"), cause=bullets("原因概要"),
                resp=bullets("対応"), meas=bullets("対策概要"), url=url, src=src)


def item_row(s, x, y, lw, vw, h, label, value, vsize=11):
    """項目別記載の1行（左=灰ラベル／右=白値・word_wrap）。x,y,… は inch。"""
    rect(s, Inches(x), Inches(y), Inches(lw), Inches(h), GRD, line=WHITE, line_w=1.0)
    add_text(s, Inches(x), Inches(y), Inches(lw), Inches(h),
             [{"runs": [(label, dict(name=JP, size=11, bold=True, color=INK))]}],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, Inches(x + lw), Inches(y), Inches(vw), Inches(h), WHITE, line=GRD, line_w=1.0)
    add_text(s, Inches(x + lw + 0.12), Inches(y), Inches(vw - 0.24), Inches(h),
             [{"runs": [(value, dict(name=JP, size=vsize, color=INK))]}],
             anchor=MSO_ANCHOR.MIDDLE)


def slide_case(prs, page, num):
    """1事例＝1スライド。上=AI再現写真3点／下=項目別記載。監修・出典・AI注記を小さく。"""
    s = blank_slide(prs)
    c = parse_case(num)
    kx = kicker(s, "重大事故事例", RED)
    add_text(s, Emu(int(kx) + int(Inches(0.25))), Inches(0.40), Inches(7.4), Inches(0.55),
             [{"runs": [(c["title"], dict(name=JP, size=18, bold=True, color=INK))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(10.0), Inches(0.40), Inches(2.72), Inches(0.55),
             [{"runs": [(c["cat"] + "／" + c["typ"], dict(name=JP, size=13, bold=True, color=RED))]}],
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    # 上段：AI再現写真3点（ラベル帯＋画像／歪みなし fit）
    photos = [("base", "現場ベース画像"), ("openai", "AI再現（OpenAI）"), ("google", "AI再現（Google）")]
    x0, gap = 0.62, 0.30
    cw = (12.08 - 2 * gap) / 3
    lt, lh, pt, ph = 1.40, 0.26, 1.68, 2.06
    for i, (key, lab) in enumerate(photos):
        lx = x0 + i * (cw + gap)
        rect(s, Inches(lx), Inches(lt), Inches(cw), Inches(lh), NAVY)
        add_text(s, Inches(lx), Inches(lt), Inches(cw), Inches(lh),
                 [{"runs": [(lab, dict(name=JP, size=10, bold=True, color=WHITE))]}],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        img = os.path.join(BASE, "photos_v16", num, key + ".png")
        place_fig(s, img, Inches(lx), Inches(pt), Inches(cw), Inches(ph), line=LGRAY)
    # 下段：項目別記載（博展テンプレ準拠・本文そのまま・文字小さめ）
    lw, vw, y = 2.05, 12.08 - 2.05, 3.96
    rows = [
        ("発生事象", _trim(c["happen"], 140), 0.66, 11),
        ("原因概要", "／".join(c["cause"][:2]), 0.54, 11),
        ("対応", "／".join(c["resp"][:2]), 0.44, 10),
        ("対策概要", "／".join(c["meas"][:3]), 0.76, 11),
    ]
    for lab, val, h, sz in rows:
        item_row(s, x0, y, lw, vw, h, lab, val, sz)
        y += h
    # 出典（小さく）＋AI再現イメージ注記（小さく）
    add_text(s, Inches(0.62), Inches(6.40), Inches(12.1), Inches(0.30),
             [{"runs": [("出典：", dict(name=JP, size=10, bold=True, color=GRAY)),
                        (c["src"] + "　" + c["url"], dict(name=JP, size=10, color=GRAY))]}])
    add_text(s, Inches(0.62), Inches(6.70), Inches(12.1), Inches(0.30),
             [{"runs": [("※ 写真はAIによる再現イメージ（実写ではありません）。公的災害事例の機序を参考にした創作（再現）事例です。",
                         dict(name=JP, size=10, color=LGRAY))]}])
    footer(s, page)  # フッター左に監修表記（金田 義太・登録第4840号）
    return page + 1


# 確定連絡先（オーナー実アドレス・捏造なし。電話番号は未確認のため載せない）
CONTACT_MAIL = "kenshi.ycc@gmail.com"


def slide_summary(prs, page):
    """まとめ＋提案：「科学的安全 × 自動化」を1枚で。価値1行・監修者名・連絡先。押し売りしない。"""
    s = blank_slide(prs)
    # 表紙と対の装丁（左ネイビー帯＋上トライカラー）
    rect(s, 0, 0, Inches(0.22), SH, NAVY)
    tricolor_rule(s, Inches(0.22), 0, Emu(int(SW) - int(Inches(0.22))), Inches(0.10))
    kicker(s, "まとめ", GREEN)
    # 中核メッセージ（大・中央）
    add_text(s, Inches(0.9), Inches(1.30), Inches(11.6), Inches(1.0),
             [{"runs": [("科学的安全", dict(name=JP, size=44, bold=True, color=NAVY)),
                        ("　×　", dict(name=JP, size=44, bold=True, color=RED)),
                        ("自動化", dict(name=JP, size=44, bold=True, color=NAVY))]}],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # 価値1行（押し売りでなく）
    add_text(s, Inches(0.9), Inches(2.42), Inches(11.6), Inches(0.5),
             [{"runs": [("経験と勘に頼らない安全教育を、速く・安く・誰でも同じ品質で。",
                         dict(name=JP, size=18, color=GRAY))]}],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # 2本柱カード（データ／AI）
    pillars = [(NAVY, "データに基づく", "約42万件の災害分析で、守る急所を科学的に特定。"),
               (RED, "AIで量産", "事故事例教材を属人化させず、速く更新し続ける。")]
    pw, ph, gap = Inches(5.70), Inches(1.78), Inches(0.40)
    x0, y0 = Inches(0.78), Inches(3.30)
    for i, (acc, head, sub) in enumerate(pillars):
        lx = Emu(int(x0) + i * (int(pw) + int(gap)))
        rect(s, lx, y0, pw, ph, PALE, line=acc, line_w=2.0)
        rect(s, lx, y0, Inches(0.16), ph, acc)
        inx = Emu(int(lx) + int(Inches(0.46)))
        inw = Emu(int(pw) - int(Inches(0.66)))
        add_text(s, inx, Emu(int(y0) + int(Inches(0.22))), inw, Inches(0.6),
                 [{"runs": [(head, dict(name=JP, size=24, bold=True, color=acc))]}])
        add_text(s, inx, Emu(int(y0) + int(Inches(0.92))), inw, Inches(0.74),
                 [{"runs": [(sub, dict(name=JP, size=14, color=INK))]}])
    # 連絡先・監修者カード（下段・控えめ）
    cy = Inches(5.62)
    rect(s, Inches(0.78), cy, Inches(11.78), Inches(0.96), NAVY)
    add_text(s, Inches(1.08), cy, Inches(7.2), Inches(0.96),
             [{"runs": [("お問い合わせ", dict(name=JP, size=12, bold=True, color=YELLOW))],
               "space_after": 2},
              {"runs": [(CONTACT_MAIL, dict(name=JP, size=17, bold=True, color=WHITE))]}],
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(8.0), cy, Inches(4.3), Inches(0.96),
             [{"runs": [("監修", dict(name=JP, size=11, bold=True, color=YELLOW))],
               "align": PP_ALIGN.RIGHT, "space_after": 2},
              {"runs": [("金田 義太", dict(name=JP, size=15, bold=True, color=WHITE))],
               "align": PP_ALIGN.RIGHT},
              {"runs": [("労働安全コンサルタント 登録第4840号",
                         dict(name=JP, size=10, color=RGBColor(0xC8, 0xD2, 0xE0)))],
               "align": PP_ALIGN.RIGHT}],
             anchor=MSO_ANCHOR.MIDDLE)
    footer(s, page)
    return page + 1


# 登録順＝スライド順（P3〜P6 でここに追記）
SLIDES = [slide_cover, slide_problem, slide_approach, slide_data1, slide_data2,
          slide_measures, slide_value, slide_samples,
          lambda prs, page: slide_case(prs, page, "N01"),
          lambda prs, page: slide_case(prs, page, "N06"),
          slide_summary]


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "proposal_hakuten.pptx")
    prs = Presentation()
    prs.slide_width = SW; prs.slide_height = SH
    page = 1
    for fn in SLIDES:
        page = fn(prs, page)
    prs.save(out)
    print("SAVED", out, "slides=", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
